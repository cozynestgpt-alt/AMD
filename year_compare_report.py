"""
전년대비 판매수수료/손익 보고서 생성 모듈

목적
- 기존 보고서 "판매수수료_연도별DB(전년대비)자동화Ver_YYYYMM.xlsx"의 구조를 개선하여 자동 생성합니다.
- 표1 시트는 월별손익DB.xlsx의 월별 손익 데이터를 기준으로 구성합니다.
- 구분 시트는 input/매장구분_행사_신규_폐점_YYYYMM.xlsx 파일을 우선 사용합니다.
- 행사/신규/폐점/동일 분류는 매월 실무자가 입력한 매장구분 파일 기준으로 반영합니다.

주의
- 판매수수료작업시트_AMD.xlsx에는 매출/수수료/공제 중심 정보가 있고,
  생산원가(V-), 영업이익(V-) 등 손익분석 컬럼은 월별손익DB 또는 매출집계_분석.xlsx 기준입니다.
- 따라서 본 보고서의 손익 지표는 월별손익DB를 기준으로 생성합니다.
"""
from __future__ import annotations

from pathlib import Path
from datetime import datetime
from collections import defaultdict
import calendar
import sys
from copy import copy

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"])
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference

NAVY = "1F3864"
BLUE = "9FD5EA"
BLUE2 = "BFE8F7"
PEACH = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "D9EAD3"
LIGHT_GRAY = "F2F2F2"
RED = "C00000"
GREEN = "1E6B3C"
YELLOW = "FFF2CC"
BORDER = Side(style="thin", color="808080")
MED = Side(style="medium", color="404040")
FMT = '#,##0'
FMT_THOUSAND = '#,##0'
PCT = '0.0%'

METRICS = ["매출금액", "수금액(V+)", "수금액(V-)", "생산원가(V-)", "영업이익(V-)", "총경비", "순이익"]
DB_MAP = {
    "매출금액": "판매금액",
    "수금액(V+)": "수금액(V+)",
    "수금액(V-)": "수금액(V-)",
    "생산원가(V-)": "생산원가(V-)",
    "영업이익(V-)": "영업이익(V-)",
    "총경비": "총경비",
    "순이익": "순이익",
}


def _month_end(year: int, month: int) -> datetime:
    return datetime(year, month, calendar.monthrange(year, month)[1])


def _to_int(v) -> int:
    if v is None or v == "":
        return 0
    try:
        return int(round(float(v)))
    except Exception:
        return 0


def _ym_to_year_month(ym: str) -> tuple[int, int]:
    ym = str(ym).strip()
    if len(ym) == 6 and ym.isdigit():
        ym = ym[:4] + "-" + ym[4:]
    dt = datetime.strptime(ym, "%Y-%m")
    return dt.year, dt.month


def _find_store_class_file(input_dir: Path, ym: str) -> Path | None:
    compact = ym.replace("-", "")
    # scoped: 연월(compact)이 포함되어 다른 달과 안 겹침 → 최상위 우선, 없으면
    #         input/이전/까지 재귀 탐색 (과거 달 재조회 지원)
    scoped = [
        f"매장구분_행사_신규_폐점_{compact}.xlsx",
        f"*매장구분*{compact}*.xlsx",
    ]
    # loose: 연월 정보 없는 느슨한 패턴 → 최상위에서만, scoped가 전부 실패했을 때 최후 수단
    loose = [
        "*매장구분*행사*신규*폐점*.xlsx",
        "*매장구분*.xlsx",
    ]
    for pat in scoped:
        found = sorted(input_dir.glob(pat))
        if found:
            return found[0]
    for pat in scoped:
        found = sorted(input_dir.rglob(pat))
        if found:
            return found[0]
    for pat in loose:
        found = sorted(input_dir.glob(pat))
        if found:
            return found[0]
    return None


def load_store_class(path: Path | None) -> dict[str, dict]:
    """{매장코드: {매장명, 구분, 비고}}"""
    if not path or not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["구분"] if "구분" in wb.sheetnames else wb.active
    result = {}
    header_row = None
    headers = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 10) + 1)]
        joined = "|".join(str(v or "") for v in vals)
        if "매장명" in joined and "구분" in joined:
            header_row = r
            for c in range(1, ws.max_column + 1):
                headers[str(ws.cell(r, c).value or "").strip()] = c
            break
    if not header_row:
        return result
    code_col = headers.get("매장코드") or headers.get("행 레이블") or 2
    name_col = headers.get("매장명") or 3
    cls_col = headers.get("구분") or 4
    note_col = headers.get("비고") or headers.get("행사명") or None
    for r in range(header_row + 1, ws.max_row + 1):
        code = str(ws.cell(r, code_col).value or "").strip()
        if not code or code.lower() == "none" or code == "총합계":
            continue
        code = code.zfill(5)
        name = str(ws.cell(r, name_col).value or "").strip()
        cls = str(ws.cell(r, cls_col).value or "").strip()
        note = str(ws.cell(r, note_col).value or "").strip() if note_col else ""
        if not cls:
            cls = "동일"
        if cls not in ("동일", "신규", "폐점", "행사"):
            # 실무 메모가 포함된 경우 기본 동일로 두되 원문은 비고에 보존
            note = (note + " " + cls).strip()
            cls = "동일"
        result[code] = {"매장코드": code, "매장명": name, "구분": cls, "비고": note}
    return result


def load_profit_db(db_path: Path) -> list[dict]:
    wb = load_workbook(db_path, read_only=True, data_only=True)
    ws = wb["년도별월별DB"] if "년도별월별DB" in wb.sheetnames else wb.active
    headers = [str(v or "").strip() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        year = row[idx.get("년", 0)]
        month = row[idx.get("월", 1)]
        code = row[idx.get("매장코드", 9)]
        name = row[idx.get("매장명", 10)]
        if not year or not month or not code:
            continue
        try:
            y = int(year)
            m = int(str(month).replace("월", ""))
        except Exception:
            continue
        rec = {
            "년": y,
            "월": m,
            "년-월": f"{y}-{m:02d}",
            "기준일": _month_end(y, m),
            "매장코드": str(code).strip().zfill(5),
            "매장명": str(name or "").strip(),
        }
        for out_name, db_name in DB_MAP.items():
            rec[out_name] = _to_int(row[idx.get(db_name, -1)] if db_name in idx else 0)
        rec["순이익율"] = _vat_excl_rate(rec["순이익"], rec["매출금액"])
        rows.append(rec)
    return rows


def _class_for(code: str, name: str, class_map: dict[str, dict], prev_sum: int, cur_sum: int) -> str:
    if code in class_map:
        return class_map[code].get("구분", "동일") or "동일"
    # 구분 파일이 없거나 누락된 매장에 대한 보조 판정
    if cur_sum and not prev_sum:
        return "신규"
    if prev_sum and not cur_sum:
        return "폐점"
    return "동일"


def _aggregate(rows: list[dict], year: int, max_month: int, class_map: dict[str, dict]) -> dict[str, dict]:
    data = defaultdict(lambda: {"매장명": "", **{m: 0 for m in METRICS}})
    for r in rows:
        if r["년"] == year and 1 <= r["월"] <= max_month:
            code = r["매장코드"]
            data[code]["매장명"] = r["매장명"]
            for m in METRICS:
                data[code][m] += _to_int(r.get(m))
    return data


def _set_title(ws, title, cell="B2", merge="B2:N2"):
    ws.merge_cells(merge)
    c = ws[cell]
    c.value = title
    c.font = Font(name="맑은 고딕", bold=True, size=16)
    c.alignment = Alignment(horizontal="center", vertical="center")


def _style_range(ws, min_row, min_col, max_row, max_col):
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _header(ws, row, cols, start_col=1, fill=BLUE):
    for i, v in enumerate(cols, start_col):
        c = ws.cell(row, i, v)
        c.font = Font(name="맑은 고딕", bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=BORDER, right=BORDER, top=MED, bottom=MED)


def _numfmt(ws, min_row, min_col, max_row, max_col, fmt=FMT_THOUSAND):
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.number_format = fmt


def _apply_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_table(ws, start_row, start_col, headers, data_rows, number_cols=None, pct_cols=None, total_row=True):
    _header(ws, start_row, headers, start_col)
    number_cols = set(number_cols or [])
    pct_cols = set(pct_cols or [])
    for r_idx, vals in enumerate(data_rows, start_row + 1):
        for c_idx, val in enumerate(vals, start_col):
            cell = ws.cell(r_idx, c_idx, val)
            cell.font = Font(name="맑은 고딕", size=9, color=RED if isinstance(val, (int, float)) and val < 0 else "000000")
            cell.alignment = Alignment(horizontal="right" if c_idx in number_cols or c_idx in pct_cols else "center", vertical="center")
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
            if c_idx in number_cols:
                cell.number_format = FMT_THOUSAND
            if c_idx in pct_cols:
                cell.number_format = PCT
    if total_row and data_rows:
        tr = start_row + len(data_rows) + 1
        for c in range(start_col, start_col + len(headers)):
            cell = ws.cell(tr, c)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="맑은 고딕", bold=True, size=9)
            cell.border = Border(left=BORDER, right=BORDER, top=MED, bottom=MED)
        ws.cell(tr, start_col, "총합계")
    return start_row + len(data_rows) + 1


def _make_display_value(v, unit=1000):
    return int(round(v / unit)) if unit else v


def _sales_vat_excl(sales_vat_incl):
    # 매출금액은 V+ 기준이므로, V- 항목(생산원가/경비/이익)과 비교할 때는 /1.1 기준을 사용합니다.
    return sales_vat_incl / 1.1 if sales_vat_incl else 0


def _cost_multiple(sales_vat_incl, production_cost_vat_excl):
    return _sales_vat_excl(sales_vat_incl) / production_cost_vat_excl if production_cost_vat_excl else 0


def _vat_excl_rate(amount_vat_excl, sales_vat_incl):
    denom = _sales_vat_excl(sales_vat_incl)
    return amount_vat_excl / denom if denom else 0


def _add_sheet_tab_color(ws, color):
    ws.sheet_properties.tabColor = color


def make_year_compare_report(ym: str, base_dir: Path) -> Path:
    year, month = _ym_to_year_month(ym)
    prev_year = year - 1
    compact = f"{year}{month:02d}"
    input_dir = base_dir / "input"
    output_dir = base_dir / "output" / f"{year}-{month:02d}"
    output_dir.mkdir(parents=True, exist_ok=True)
    # 경영보고서 폴더에도 복사하기 쉽도록 루트 output에도 생성
    db_path = base_dir / "DB" / "월별손익DB.xlsx"
    class_path = _find_store_class_file(input_dir, f"{year}-{month:02d}")

    rows = load_profit_db(db_path)
    class_map = load_store_class(class_path)
    prev = _aggregate(rows, prev_year, month, class_map)
    cur = _aggregate(rows, year, month, class_map)
    all_codes = sorted(set(prev) | set(cur) | set(class_map), key=lambda x: x)

    # 구분 부여
    cls_by_code = {}
    names = {}
    for code in all_codes:
        name = (class_map.get(code, {}).get("매장명") or cur.get(code, {}).get("매장명") or prev.get(code, {}).get("매장명") or "")
        names[code] = name
        cls_by_code[code] = _class_for(code, name, class_map, prev.get(code, {}).get("매출금액", 0), cur.get(code, {}).get("매출금액", 0))

    wb = Workbook()
    # 기본 시트 제거
    wb.remove(wb.active)

    # 1. 표1
    ws = wb.create_sheet("표1")
    _add_sheet_tab_color(ws, NAVY)
    headers = ["년", "월", "년-월", "매장코드", "매장명", "구분", "매출금액", "수금액(V+)", "수금액(V-)", "생산원가(V-)", "영업이익(V-)", "총경비", "영업이익-총경비\n(순이익)", "영업이익-총경비\n(순이익율)"]
    _header(ws, 1, headers, 1, NAVY)
    for cell in ws[1]:
        cell.font = Font(name="맑은 고딕", bold=True, color=WHITE, size=10)
    out_rows = []
    for r in rows:
        if r["년"] in (prev_year, year) and 1 <= r["월"] <= month:
            code = r["매장코드"]
            out_rows.append([
                r["년"], f"{r['월']:02d}", r["년-월"], code, r["매장명"], cls_by_code.get(code, "동일"),
                r["매출금액"], r["수금액(V+)"], r["수금액(V-)"], r["생산원가(V-)"],
                r["영업이익(V-)"], r["총경비"], r["순이익"], r["순이익율"]
            ])
    out_rows.sort(key=lambda x: (x[0], x[1], x[3]))
    for i, vals in enumerate(out_rows, 2):
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name="맑은 고딕", size=9, color=RED if isinstance(v, (int, float)) and v < 0 else "000000")
            c.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
            c.alignment = Alignment(horizontal="right" if j >= 7 else "center", vertical="center")
            if 7 <= j <= 13:
                c.number_format = FMT_THOUSAND
            if j == 14:
                c.number_format = PCT
    _apply_widths(ws, [7,6,10,10,20,10,13,13,13,13,13,13,15,13])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:N{len(out_rows)+1}"

    # 2. 구분
    ws = wb.create_sheet("구분")
    _add_sheet_tab_color(ws, GREEN)
    ws["B1"] = f"매장구분 기준표  ▶  {year}-{month:02d}"
    ws["B1"].font = Font(name="맑은 고딕", bold=True, size=14)
    ws["B2"] = f"기준 파일: {class_path.name if class_path else '없음 - DB 기준 자동 분류'}"
    ws["B2"].font = Font(name="맑은 고딕", size=9, color="666666")
    headers_cls = ["매장코드", "매장명", "구분", "동일_순번", "폐점_순번", "행사_순번", "신규_순번", "비고"]
    _header(ws, 3, headers_cls, 2, BLUE)
    order_counter = defaultdict(int)
    cls_order = {"동일": 1, "폐점": 2, "행사": 3, "신규": 4}
    for r_idx, code in enumerate(sorted(all_codes, key=lambda c: (cls_order.get(cls_by_code.get(c, "동일"), 9), c)), 4):
        cls = cls_by_code.get(code, "동일")
        order_counter[cls] += 1
        vals = [code, names.get(code,""), cls, None, None, None, None, class_map.get(code, {}).get("비고", "")]
        if cls == "동일": vals[3] = order_counter[cls]
        if cls == "폐점": vals[4] = order_counter[cls]
        if cls == "행사": vals[5] = order_counter[cls]
        if cls == "신규": vals[6] = order_counter[cls]
        for c_idx, v in enumerate(vals, 2):
            cell = ws.cell(r_idx, c_idx, v)
            cell.font = Font(name="맑은 고딕", size=9)
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
            cell.alignment = Alignment(horizontal="center" if c_idx != 3 else "left", vertical="center")
    _apply_widths(ws, [3,10,22,10,10,10,10,10,24])
    ws.freeze_panes = "B4"

    def _metric(code, data, metric):
        return data.get(code, {}).get(metric, 0)

    # 3. 피벗값_구분
    ws = wb.create_sheet("피벗값_구분")
    _add_sheet_tab_color(ws, NAVY)
    _set_title(ws, "전년동기대비 증감액 및 증감율", "B2", "B2:N2")
    ws["L3"] = "단위: 천원"
    ws["L3"].font = Font(name="맑은 고딕", size=9)
    groups = ["동일", "신규", "폐점", "행사"]
    row_data = []
    totals_by_cls = {}
    for cls in groups:
        codes = [c for c in all_codes if cls_by_code.get(c) == cls]
        p = {m: sum(_metric(c, prev, m) for c in codes) for m in METRICS}
        q = {m: sum(_metric(c, cur, m) for c in codes) for m in METRICS}
        totals_by_cls[cls] = (p, q)
        row_data.append([cls, _make_display_value(p["매출금액"]), _make_display_value(q["매출금액"]), _make_display_value(p["수금액(V+)"]), _make_display_value(q["수금액(V+)"]), _make_display_value(p["생산원가(V-)"]), _make_display_value(q["생산원가(V-)"]), _make_display_value(p["총경비"]), _make_display_value(q["총경비"]), _make_display_value(p["영업이익(V-)"]), _make_display_value(q["영업이익(V-)"])])
    # total
    total_prev = {m: sum(totals_by_cls[g][0][m] for g in groups) for m in METRICS}
    total_cur = {m: sum(totals_by_cls[g][1][m] for g in groups) for m in METRICS}
    row_data.append(["총합계", _make_display_value(total_prev["매출금액"]), _make_display_value(total_cur["매출금액"]), _make_display_value(total_prev["수금액(V+)"]), _make_display_value(total_cur["수금액(V+)"]), _make_display_value(total_prev["생산원가(V-)"]), _make_display_value(total_cur["생산원가(V-)"]), _make_display_value(total_prev["총경비"]), _make_display_value(total_cur["총경비"]), _make_display_value(total_prev["영업이익(V-)"]), _make_display_value(total_cur["영업이익(V-)"])])
    ws.merge_cells("B4:C4"); ws["B4"] = "매출금액"
    ws.merge_cells("D4:E4"); ws["D4"] = "수금액(V+)"
    ws.merge_cells("F4:G4"); ws["F4"] = "생산원가(V-)"
    ws.merge_cells("H4:I4"); ws["H4"] = "총경비"
    ws.merge_cells("J4:K4"); ws["J4"] = "영업이익(V-)"
    for c in ["B4","D4","F4","H4","J4"]:
        ws[c].fill = PatternFill("solid", fgColor=BLUE)
        ws[c].font = Font(name="맑은 고딕", bold=True)
        ws[c].alignment = Alignment(horizontal="center")
    _write_table(ws, 5, 1, ["구분", prev_year, year, prev_year, year, prev_year, year, prev_year, year, prev_year, year], row_data, number_cols=set(range(2,12)), total_row=False)
    # 증감률 표
    start = 13
    ws["B12"] = f"(1~{month}월) 전기대비 매출 증가분의 당기 매출기여도"
    ws["B12"].font = Font(name="맑은 고딕", bold=True)
    diff_rows = []
    for cls in groups + ["총합계"]:
        p, q = (total_prev, total_cur) if cls == "총합계" else totals_by_cls[cls]
        diff_rows.append([
            cls,
            _make_display_value(q["매출금액"]-p["매출금액"]), (q["매출금액"]-p["매출금액"])/total_cur["매출금액"] if total_cur["매출금액"] else 0,
            _make_display_value(q["수금액(V+)"]-p["수금액(V+)"]), (q["수금액(V+)"]-p["수금액(V+)"])/total_cur["수금액(V+)"] if total_cur["수금액(V+)"] else 0,
            _make_display_value(q["생산원가(V-)"]-p["생산원가(V-)"]), (q["생산원가(V-)"]-p["생산원가(V-)"])/total_cur["생산원가(V-)"] if total_cur["생산원가(V-)"] else 0,
            _make_display_value(q["총경비"]-p["총경비"]), (q["총경비"]-p["총경비"])/total_cur["총경비"] if total_cur["총경비"] else 0,
            _make_display_value(q["영업이익(V-)"]-p["영업이익(V-)"]), (q["영업이익(V-)"]-p["영업이익(V-)"])/total_cur["영업이익(V-)"] if total_cur["영업이익(V-)"] else 0,
        ])
    _write_table(ws, start, 1, ["구분","매출차이","매출기여도","수금액(V+)차이","수금기여도","생산원가(V-)차이","원가기여도","총경비차이","경비기여도","영업이익(V-)차이","이익기여도"], diff_rows, number_cols={2,4,6,8,10}, pct_cols={3,5,7,9,11}, total_row=False)
    _apply_widths(ws, [12,13,13,13,13,13,13,13,13,13,13,14,14])

    # Helper for store rows
    def _store_rows_for(cls: str):
        codes = [c for c in all_codes if cls_by_code.get(c) == cls]
        rows2 = []
        for c in codes:
            p = prev.get(c, {m: 0 for m in METRICS})
            q = cur.get(c, {m: 0 for m in METRICS})
            rows2.append((c, names.get(c,""), p, q))
        return rows2

    # 4. 동일매장분석
    ws = wb.create_sheet("피벗값_매장명(동일)")
    _add_sheet_tab_color(ws, GREEN)
    _set_title(ws, "동일매장분석", "B2", "B2:V2")
    ws["U3"] = "단위: 천원"
    same_rows = []
    for code, name, p, q in _store_rows_for("동일"):
        sales_growth = (q["매출금액"] - p["매출금액"]) / p["매출금액"] if p["매출금액"] else 0
        collect_rate = q["수금액(V+)"] / q["매출금액"] if q["매출금액"] else 0
        cost_mult = _cost_multiple(q["매출금액"], q["생산원가(V-)"])
        exp_rate = _vat_excl_rate(q["총경비"], q["매출금액"])
        profit_rate = _vat_excl_rate(q["영업이익(V-)"], q["매출금액"])
        diff_profit = q["영업이익(V-)"] - p["영업이익(V-)"]
        same_rows.append([code, name, _make_display_value(p["매출금액"]), _make_display_value(q["매출금액"]), sales_growth, _make_display_value(p["수금액(V+)"]), _make_display_value(q["수금액(V+)"]), collect_rate, _make_display_value(p["생산원가(V-)"]), _make_display_value(q["생산원가(V-)"]), cost_mult, _make_display_value(p["총경비"]), _make_display_value(q["총경비"]), exp_rate, _make_display_value(p["영업이익(V-)"]), q["영업이익(V-)"], profit_rate, 0, _make_display_value(diff_profit)])
    same_rows.sort(key=lambda x: x[0])   # A열 매장코드 오름차순
    # 공헌이익률 = 해당 점포 당해년도 영업이익 / 동일매장 전체 영업이익 합계 (R열)
    _total_cur_profit = sum(r[15] for r in same_rows)  # index15 = 당해년도 영업이익 raw값
    for r in same_rows:
        r[17] = r[15] / _total_cur_profit if _total_cur_profit else 0  # index17 = 공헌이익률
    # P열(index15)을 표시용 천원 단위로 변환 (raw→display, 공헌이익률 계산 후)
    for r in same_rows:
        r[15] = _make_display_value(r[15])
    headers_same = ["동일", "매장명", f"{prev_year}\n매출", f"{year}\n매출", "당기기준\n신장률", f"{prev_year}\n수금", f"{year}\n수금", "매출대비\n수금율", f"{prev_year}\n생산원가", f"{year}\n생산원가", "매출대비\n원가배수", f"{prev_year}\n총경비", f"{year}\n총경비", "매출대비\n경비율", f"{prev_year}\n영업이익", f"{year}\n영업이익", "매출대비\n이익률", "공헌\n이익률", "증감액"]
    _write_table(ws, 5, 1, headers_same, same_rows, number_cols={3,4,6,7,9,10,12,13,15,16,19}, pct_cols={5,8,14,17,18}, total_row=False)
    _apply_widths(ws, [8,20,11,11,9,11,11,9,11,11,9,11,11,9,11,11,9,9,11])
    ws.freeze_panes = "A6"

    # 5. 동일매장 영업이익순
    ws2 = wb.create_sheet("피벗값_매장명(동일) 값_영업이익순")
    _add_sheet_tab_color(ws2, GREEN)
    _set_title(ws2, "동일매장 영업이익 높은순", "B2", "B2:V2")
    same_profit_rows = sorted(same_rows, key=lambda x: x[15], reverse=True)
    _write_table(ws2, 5, 1, headers_same, same_profit_rows, number_cols={3,4,6,7,9,10,12,13,15,16,19}, pct_cols={5,8,14,17,18}, total_row=False)
    _apply_widths(ws2, [8,20,11,11,9,11,11,9,11,11,9,11,11,9,11,11,9,9,11])
    ws2.freeze_panes = "A6"

    # event/closed/new sheets
    def _simple_class_sheet(sheet_name, title, cls, include_both=True):
        ws = wb.create_sheet(sheet_name)
        _add_sheet_tab_color(ws, PEACH if cls in ("행사", "신규") else RED)
        _set_title(ws, title, "B3", "B3:N3")
        ws["M4"] = "단위: 천원"
        rows3 = []
        for code, name, p, q in _store_rows_for(cls):
            note = class_map.get(code, {}).get("비고", "")
            if include_both:
                rows3.append([note, code, name, _make_display_value(p["매출금액"]), _make_display_value(q["매출금액"]), _make_display_value(p["수금액(V+)"]), _make_display_value(q["수금액(V+)"]), _make_display_value(p["생산원가(V-)"]), _make_display_value(q["생산원가(V-)"]), _make_display_value(p["총경비"]), _make_display_value(q["총경비"]), _make_display_value(p["영업이익(V-)"]), _make_display_value(q["영업이익(V-)"])])
            else:
                src = p if cls == "폐점" else q
                rows3.append([note, code, name, _make_display_value(src["매출금액"]), _make_display_value(src["수금액(V+)"]), _make_display_value(src["생산원가(V-)"]), _make_display_value(src["총경비"]), _make_display_value(src["영업이익(V-)"])])
        if include_both:
            hdr = ["비고", cls, "매장명", f"{prev_year}\n매출", f"{year}\n매출", f"{prev_year}\n수금", f"{year}\n수금", f"{prev_year}\n생산원가", f"{year}\n생산원가", f"{prev_year}\n총경비", f"{year}\n총경비", f"{prev_year}\n영업이익", f"{year}\n영업이익"]
            _write_table(ws, 6, 1, hdr, rows3, number_cols=set(range(4,14)), total_row=False)
            _apply_widths(ws, [16,8,20,11,11,11,11,11,11,11,11,11,11])
        else:
            period_label = f"{prev_year}년(1~{month}월)총 합계" if cls == "폐점" else f"{year}년(1~{month}월)총 합계"
            ws["C5"] = period_label
            ws["C5"].font = Font(name="맑은 고딕", bold=True)
            hdr = ["비고", cls, "매장명", "매출금액", "수금액(V+)", "생산원가(V-)", "총경비", "영업이익(V-)"]
            _write_table(ws, 6, 1, hdr, rows3, number_cols=set(range(4,9)), total_row=False)
            _apply_widths(ws, [16,8,20,12,12,12,12,12])
        ws.freeze_panes = "A7"
        return ws

    _simple_class_sheet("피벗값_매장명 (폐점)", "폐점매장분석", "폐점", include_both=False)
    _simple_class_sheet("피벗값_매장명 (행사)", "행사매장분석", "행사", include_both=True)
    _simple_class_sheet("피벗값_매장명 (신규)", "신규매장분석", "신규", include_both=False)

    # Dashboard
    ws = wb.create_sheet("Dashboard", 0)
    _add_sheet_tab_color(ws, NAVY)
    ws.merge_cells("A1:H1")
    ws["A1"] = f"판매수수료 전년대비 보고서 Dashboard  ▶  {year}년 1~{month}월"
    ws["A1"].font = Font(name="맑은 고딕", bold=True, size=16, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")
    kpis = [
        ("매출금액", total_prev["매출금액"], total_cur["매출금액"]),
        ("수금액(V+)", total_prev["수금액(V+)"], total_cur["수금액(V+)"]),
        ("생산원가(V-)", total_prev["생산원가(V-)"], total_cur["생산원가(V-)"]),
        ("총경비", total_prev["총경비"], total_cur["총경비"]),
        ("영업이익(V-)", total_prev["영업이익(V-)"], total_cur["영업이익(V-)"]),
        ("순이익", total_prev["순이익"], total_cur["순이익"]),
    ]
    _write_table(ws, 3, 1, ["항목", prev_year, year, "증감액", "증감률"], [[n, _make_display_value(p), _make_display_value(q), _make_display_value(q-p), (q-p)/p if p else 0] for n,p,q in kpis], number_cols={2,3,4}, pct_cols={5}, total_row=False)
    ws["A12"] = "분류별 요약"
    ws["A12"].font = Font(name="맑은 고딕", bold=True, size=12)
    _write_table(ws, 13, 1, ["구분", f"{prev_year} 매출", f"{year} 매출", "매출증감", f"{prev_year} 영업이익", f"{year} 영업이익", "이익증감"], [[g, _make_display_value(totals_by_cls[g][0]["매출금액"]), _make_display_value(totals_by_cls[g][1]["매출금액"]), _make_display_value(totals_by_cls[g][1]["매출금액"]-totals_by_cls[g][0]["매출금액"]), _make_display_value(totals_by_cls[g][0]["영업이익(V-)"]), _make_display_value(totals_by_cls[g][1]["영업이익(V-)"]), _make_display_value(totals_by_cls[g][1]["영업이익(V-)"]-totals_by_cls[g][0]["영업이익(V-)"])] for g in groups], number_cols={2,3,4,5,6,7}, total_row=False)
    _apply_widths(ws, [16,14,14,14,12,14,14,14])

    # 차트: KPI 증감
    try:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "전년대비 주요 지표"
        chart.y_axis.title = "천원"
        data = Reference(ws, min_col=2, max_col=3, min_row=3, max_row=9)
        cats = Reference(ws, min_col=1, min_row=4, max_row=9)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 8
        chart.width = 18
        ws.add_chart(chart, "G3")
    except Exception:
        pass

    # workbook properties and save
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    align = copy(cell.alignment)
                    align.vertical = "center"
                    cell.alignment = align
        ws.freeze_panes = ws.freeze_panes or "A2"

    out_path = output_dir / "판매수수료_전년대비보고서.xlsx"
    wb.save(out_path)
    # also copy to root output for quick access
    root_out = base_dir / "output" / "판매수수료_전년대비보고서.xlsx"
    try:
        import shutil
        shutil.copy2(out_path, root_out)
    except Exception:
        pass
    print(f"     ✅ 판매수수료_전년대비보고서.xlsx 생성: {out_path}")
    return out_path



def run(ym: str, base_dir: Path | str = None) -> Path:
    base = Path(base_dir) if base_dir else Path(__file__).parent
    return make_year_compare_report(ym, base)



def _find_latest_output_month(base_dir: Path) -> str:
    """기준월을 자동으로 찾습니다.

    1순위: output/YYYY-MM 형식의 최신 월 폴더
    2순위: DB/월별손익DB.xlsx의 최신 년-월
    """
    output_dir = base_dir / "output"
    months = []

    if output_dir.exists():
        for folder in output_dir.iterdir():
            if not folder.is_dir():
                continue
            try:
                datetime.strptime(folder.name, "%Y-%m")
                months.append(folder.name)
            except ValueError:
                continue

    if months:
        return sorted(months)[-1]

    # output 월별 폴더가 없으면 월별손익DB에서 최신 년-월을 찾습니다.
    db_path = base_dir / "DB" / "월별손익DB.xlsx"
    if not db_path.exists():
        raise FileNotFoundError(f"기준월을 찾을 수 없습니다. output 월별 폴더와 DB 파일을 확인하세요: {db_path}")

    wb = load_workbook(db_path, read_only=True, data_only=True)
    ws = wb["년도별월별DB"] if "년도별월별DB" in wb.sheetnames else wb.active
    headers = [str(v or "").strip() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    ym_idx = headers.index("년-월") if "년-월" in headers else None
    y_idx = headers.index("년") if "년" in headers else None
    m_idx = headers.index("월") if "월" in headers else None

    for row in ws.iter_rows(min_row=2, values_only=True):
        ym = None
        if ym_idx is not None:
            ym = row[ym_idx]
        elif y_idx is not None and m_idx is not None:
            y, m = row[y_idx], row[m_idx]
            if y and m:
                ym = f"{int(y)}-{int(str(m).replace('월','')):02d}"
        if not ym:
            continue
        ym = str(ym).strip()
        try:
            datetime.strptime(ym, "%Y-%m")
            months.append(ym)
        except ValueError:
            continue

    if not months:
        raise FileNotFoundError("output 폴더와 월별손익DB에서 YYYY-MM 형식의 기준월을 찾지 못했습니다.")

    return sorted(set(months))[-1]


if __name__ == "__main__":
    base = Path(__file__).parent

    if len(sys.argv) >= 2:
        ym = sys.argv[1].strip()
        if len(ym) == 6 and ym.isdigit():
            ym = ym[:4] + "-" + ym[4:]
    else:
        ym = _find_latest_output_month(base)

    print()
    print("=" * 60)
    print(" 전년대비보고서 생성")
    print("=" * 60)
    print(f" 기준월: {ym}")
    print()
    run(ym, base)
    print()
    print("전년대비보고서 생성 완료")
