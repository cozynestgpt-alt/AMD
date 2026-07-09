# -*- coding: utf-8 -*-
"""
AMD 공통 모듈
- 색상 팔레트
- 스타일 함수 (st, bdr, set_col_widths)
- STORE_CODE_MAP / CODE_TO_STORE / ONLINE_CODE_MAP

각 모듈에서:
    from common import NAVY, BLUE, st, bdr, STORE_CODE_MAP
"""

from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ══════════════════════════════════════════════════════════
# 입력 파일 탐색 (input/YYYY-MM/ 폴더 전용)
# ══════════════════════════════════════════════════════════
# 2026-07부터 input/ 은 매달 input/YYYY-MM/ 폴더 하나에만 그 달 자료가 들어있는
# 구조로 바뀌었다 (예전 input/이전/ 재귀 탐색 방식 폐지). base로 넘긴 폴더 안에서만
# 찾으므로 다른 달 파일과 섞일 걱정이 없다.
def find_input(scoped_patterns, loose_patterns=(), base: Path = None):
    """base(정산월 폴더) 안에서 scoped_patterns를 먼저 시도하고,
    전부 실패하면 loose_patterns를 시도한다. 반환: 찾은 Path 리스트(중복 제거, 우선순위 순)."""
    if base is None or not base.exists():
        return []
    seen, out = set(), []
    def _add(paths):
        for p in paths:
            key = str(p)
            if key not in seen:
                seen.add(key); out.append(p)
    for pat in scoped_patterns:
        _add(sorted(base.glob(pat)))
    if out:
        return out
    for pat in loose_patterns:
        _add(sorted(base.glob(pat)))
    return out

def find_input_one(scoped_patterns, loose_patterns=(), base: Path = None):
    found = find_input(scoped_patterns, loose_patterns, base)
    return found[0] if found else None

# 정산월 폴더(input/YYYY-MM/) 안에 있어야 하는 파일 종류.
# required=True 인 파일이 없으면 run.py가 무엇이 빠졌는지 명확히 나열하고 중단한다
# (매출 0 등으로 조용히 대체하지 않음). required=False 파일은 없으면 경고만 하고
# 해당 항목 0/건너뜀으로 처리해 왔던 기존 동작을 유지한다.
INPUT_FILE_SPECS = [
    ("사원마스터",                              ["사원마스터*.xlsx"],                                True),
    ("영판매(매출)",                            ["영*판매*.xlsx"],                                   True),
    ("경비내역서",                              ["★*경비내역서*.xlsx", "*경비내역서*.xlsx"],          False),
    ("매장공제건집계",                          ["*매장공제건*.xlsx", "◈*.xlsx"],                    False),
    ("재고실사공제건합계",                      ["*재고실사*공제건*합계*.xlsx", "*재고실사*.xlsx"],   False),
    ("사원현황(아르바이트)",                    ["사원현황*.xlsx"],                                  False),
    ("급여상여명세서_일용직",                   ["급여상여명세서*일용직*.xlsx"],                      False),
    ("급여상여명세서_매장직",                   ["급여상여명세서*매장직*.xlsx"],                      False),
    ("월별근태(아르바이트)",                    ["월별근태*.xlsx"],                                  False),
    ("매장전화요금내역",                        ["매장전화요금내역*.xlsx", "*전화요금*.xlsx"],        False),
    ("중간관리매장_매입세금계산서_기재사항정보", ["중간관리매장_매입세금계산서*.xlsx"],               False),
    ("로젠고객직배",                            ["*로젠*고객직배*.xlsx", "*고객직배*.xlsx"],          False),
    ("매장구분_행사_신규_폐점",                 ["매장구분_행사_신규_폐점*.xlsx"],                    False),
]

def check_input_month(month_dir: Path):
    """month_dir(input/YYYY-MM/)을 점검해 (필수 누락 라벨 목록, 선택 누락 라벨 목록)을 반환."""
    missing_required, missing_optional = [], []
    for label, patterns, required in INPUT_FILE_SPECS:
        if not find_input_one(patterns, base=month_dir):
            (missing_required if required else missing_optional).append(label)
    return missing_required, missing_optional

# ══════════════════════════════════════════════════════════
# 색상 팔레트
# ══════════════════════════════════════════════════════════
NAVY        = "1F3864"
BLUE        = "2E5EAA"
LIGHT_BLUE  = "D9E2F3"
PALE_BLUE   = "EEF3FB"
GREEN       = "1E6B3C"
LIGHT_GREEN = "E2EFDA"
AMBER       = "FFF2CC"
AMBER_DARK  = "BF8F00"
WHITE       = "FFFFFF"
GRAY        = "F2F2F2"
GRAY_MID    = "D6D6D6"
RED_LIGHT   = "FFE0E0"
CORAL       = "C00000"
INPUT_BLUE  = "0000FF"
LINK_GREEN  = "008000"
AUTO_BLACK  = "000000"

FMT_COMMA = '#,##0'
FMT_PCT   = '0.0%'

# ══════════════════════════════════════════════════════════
# 공통 스타일 함수
# ══════════════════════════════════════════════════════════
def st(ws, row, col, val=None, bg=None, fg="000000", bold=False,
       size=9, align="center", wrap=False, fmt=None, italic=False):
    """셀 값+스타일 한번에 설정"""
    c = ws.cell(row=row, column=col)
    if val is not None:
        c.value = val
    c.font      = Font(name="맑은 고딕", bold=bold, color=fg, size=size, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if fmt:
        c.number_format = fmt
    return c

def bdr(ws, r1, c1, r2, c2, outer="medium", inner="thin"):
    """범위 테두리 설정"""
    tk = Side(style=outer)
    tn = Side(style=inner)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = Border(
                left   = tk if c == c1 else tn,
                right  = tk if c == c2 else tn,
                top    = tk if r == r1 else tn,
                bottom = tk if r == r2 else tn,
            )

# run.py에서 사용하는 이름 (borders)도 동일하게 제공
borders = bdr

def set_col_widths(ws, widths):
    """열 너비 일괄 설정 (1번 열부터 순서대로)"""
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ══════════════════════════════════════════════════════════
# 매장 코드 매핑 (단일 관리)
# ══════════════════════════════════════════════════════════
STORE_CODE_MAP = {
    "롯데잠실점":              "00101",
    "롯데관악점":              "00102",
    "롯데광주점":              "00106",
    "롯데강남점":              "00110",
    "롯데포항점":              "00111",
    "롯데울산점":              "00112",
    "롯데동래점":              "00113",
    "롯데창원점":              "00114",
    "롯데상인점":              "00119",
    "롯데미아점":              "00120",
    "롯데구리점":              "00125",
    "롯데영등포점":            "00126",
    "롯데부산본점":            "00128",
    "롯데평촌점":              "00129",
    "롯데안산점":              "00130",
    "롯데노원점":              "00131",
    "롯데광복점":              "00132",
    "롯데수원점":              "00135",
    "롯데대구점":              "00137",
    "롯데전주점":              "00139",
    "롯데인천터미널점":        "00141",
    "롯데동탄점":              "00142",
    "롯데잠실캐슬프라자점":    "00143",
    "현대천호점":              "00201",
    "현대울산점":              "00203",
    "현대압구정본점":          "00207",
    "현대미아점":              "00208",
    "현대중동점":              "00210",
    "현대킨텍스점":            "00212",
    "현대대구점":              "00213",
    "현대충청점":              "00214",
    "현대동구점":              "00218",
    "현대판교점":              "00219",
    "현대신촌점":              "00222",
    "신세계광주점":            "00301",
    "신세계마산점":            "00303",
    "신세계강남점":            "00304",
    "신세계경기점":            "00306",
    "신세계센텀점":            "00307",
    "신세계타임스퀘어점":      "00308",
    "신세계천안아산점":        "00309",
    "신세계의정부점":          "00310",
    "신세계김해점":            "00311",
    "신세계스타필드하남점":    "00313",
    "신세계대구점":            "00314",
    "신세계대전점":            "00315",
    "갤러리아천안점":          "00503",
    "갤러리아광교점":          "00504",
    "갤러리아진주점":          "00505",
    "갤러리아타임월드점":      "00506",
    "AK프라자분당점":          "00601",
    "AK프라자수원점":          "00603",
    "신세계본점":              "01305",
    "롯데아울렛김해점":        "05002",
    "롯데아울렛이천점":        "05003",
    "롯데아울렛광명점":        "05004",
    "롯데아울렛고양점":        "05005",
    "롯데아울렛광교점":        "05006",
    "롯데아울렛군산점":        "05007",
    "롯데아울렛동부산점":      "05009",
    "롯데아울렛서울역점":      "05014",
    "LF스퀘어양주점":          "05201",
    "신세계아울렛파주점":      "05552",
    "현대아울렛대전점":        "05601",
    "현대아울렛남양주점":      "05602",
}

# 코드 → 매장명 역매핑
CODE_TO_STORE = {v: k for k, v in STORE_CODE_MAP.items()}

# 온라인 코드 → 오프라인 매장명
ONLINE_CODE_MAP = {
    "12101": "롯데잠실점",       "12102": "롯데관악점",      "12106": "롯데광주점",
    "12110": "롯데강남점",       "12111": "롯데포항점",      "12112": "롯데울산점",
    "12113": "롯데동래점",       "12114": "롯데창원점",      "12119": "롯데상인점",
    "12120": "롯데미아점",       "12125": "롯데구리점",      "12126": "롯데영등포점",
    "12128": "롯데부산본점",     "12129": "롯데평촌점",      "12130": "롯데안산점",
    "12131": "롯데노원점",       "12132": "롯데광복점",      "12135": "롯데수원점",
    "12137": "롯데대구점",       "12139": "롯데전주점",      "12141": "롯데인천터미널점",
    "12142": "롯데동탄점",       "12143": "롯데잠실캐슬프라자점",
    "12201": "현대천호점",       "12203": "현대울산점",      "12207": "현대압구정본점",
    "12208": "현대미아점",       "12210": "현대중동점",      "12212": "현대킨텍스점",
    "12213": "현대대구점",       "12214": "현대충청점",      "12218": "현대동구점",
    "12219": "현대판교점",       "12222": "현대신촌점",
    "12301": "신세계광주점",     "12303": "신세계마산점",    "12304": "신세계강남점",
    "12306": "신세계경기점",     "12307": "신세계센텀점",    "12308": "신세계타임스퀘어점",
    "12309": "신세계천안아산점", "12310": "신세계의정부점",  "12311": "신세계김해점",
    "12313": "신세계스타필드하남점", "12314": "신세계대구점","12315": "신세계대전점",
    "12503": "갤러리아천안점",   "12504": "갤러리아광교점",  "12505": "갤러리아진주점",
    "12506": "갤러리아타임월드점",
    "12601": "AK프라자분당점",   "12603": "AK프라자수원점",
    "13305": "신세계본점",
    "15002": "롯데아울렛김해점", "15003": "롯데아울렛이천점",
    "15004": "롯데아울렛광명점", "15005": "롯데아울렛고양점",
    "15006": "롯데아울렛광교점", "15007": "롯데아울렛군산점",
    "15009": "롯데아울렛동부산점","15201": "LF스퀘어양주점",
    "15552": "신세계아울렛파주점",
    "15601": "현대아울렛대전점", "15602": "현대아울렛남양주점",
}
