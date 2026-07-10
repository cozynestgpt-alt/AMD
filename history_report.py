# -*- coding: utf-8 -*-
"""월별손익DB 기준 경영진용 연도/분기/시즌 분석 보고서 생성 모듈 v3"""
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

NAVY='1F3864'; BLUE='2E5EAA'; GREEN='1E6B3C'; ORANGE='F4B942'; WHITE='FFFFFF'; PALE='EEF3FB'; GRAY='F2F2F2'; RED='C00000'
FMT='#,##0'; PCT='0.00%'
METRICS=['판매금액','수금액(V+)','수금액(V-)','생산원가(V-)','영업이익(V-)','총경비','순이익']

def _n(v):
    if v is None or v == '': return 0
    if isinstance(v,(int,float)): return v
    try: return float(str(v).replace(',','').replace('%',''))
    except: return 0

def st(ws,r,c,v=None,bg=None,fg='000000',bold=False,size=9,align='center',fmt=None):
    cell=ws.cell(r,c,v)
    cell.font=Font(name='맑은 고딕',bold=bold,color=fg,size=size)
    cell.alignment=Alignment(horizontal=align,vertical='center')
    if bg: cell.fill=PatternFill('solid',fgColor=bg)
    if fmt: cell.number_format=fmt
    return cell

def header(ws, row, values, bg=NAVY):
    for i,v in enumerate(values,1): st(ws,row,i,v,bg=bg,fg=WHITE,bold=True,size=9)

def format_sheet(ws):
    thin=Side(style='thin',color='D9D9D9')
    for row in ws.iter_rows():
        for c in row: c.border=Border(bottom=thin)
    ws.freeze_panes='A2'
    ws.sheet_view.showGridLines=False

def load_db(db_path: Path):
    wb=load_workbook(db_path, data_only=True, read_only=True)
    ws=wb['년도별월별DB']
    headers=[str(c.value).strip() for c in ws[1]]
    rows=[]
    for vals in ws.iter_rows(min_row=2, values_only=True):
        if not any(vals): continue
        rows.append({headers[i]: vals[i] for i in range(min(len(headers),len(vals)))})
    return rows

def _sales_vat_excl(sales_vat_incl):
    return sales_vat_incl / 1.1 if sales_vat_incl else 0


def _vat_excl_rate(amount_vat_excl, sales_vat_incl):
    denom = _sales_vat_excl(sales_vat_incl)
    return amount_vat_excl / denom if denom else 0


def aggregate(rows, keys):
    out=defaultdict(lambda: {m:0 for m in METRICS})
    for r in rows:
        k=tuple(r.get(x) for x in keys)
        for m in METRICS: out[k][m]+=_n(r.get(m))
    result=[]
    for k,v in out.items():
        sales=v['판매금액']; profit=v['순이익']
        row={keys[i]:k[i] for i in range(len(keys))}
        row.update(v); row['순이익율']=_vat_excl_rate(profit, sales)
        result.append(row)
    result.sort(key=lambda x: tuple(str(x.get(k)) for k in keys))
    return result

def write_table(ws, rows, cols):
    header(ws,1,cols)
    for ri,r in enumerate(rows,2):
        for ci,col in enumerate(cols,1):
            v=r.get(col)
            fmt=PCT if '율' in col else (FMT if col in METRICS else None)
            align='left' if col in ('매장명','시즌명') else 'center'
            st(ws,ri,ci,v,fmt=fmt,align=align,size=9)
    for ci,col in enumerate(cols,1):
        ws.column_dimensions[get_column_letter(ci)].width = 22 if col in ('매장명','시즌명') else 14
    format_sheet(ws)

def make_report(db_path: Path, out_path: Path):
    rows=load_db(db_path)
    wb=Workbook()
    # Dashboard
    ws=wb.active; ws.title='Dashboard'
    latest=sorted({r['년-월'] for r in rows})[-1] if rows else ''
    latest_rows=[r for r in rows if r.get('년-월')==latest]
    total=aggregate(latest_rows,[])[0] if latest_rows else {m:0 for m in METRICS}|{'순이익율':0}
    ws.merge_cells('A1:H1'); st(ws,1,1,f'경영 Dashboard - 최신월 {latest}',bg=NAVY,fg=WHITE,bold=True,size=15)
    kpis=[('매출액',total['판매금액']),('영업이익',total['영업이익(V-)']),('총경비',total['총경비']),('순이익',total['순이익']),('순이익율',total['순이익율'])]
    for i,(name,val) in enumerate(kpis):
        c=1+i*2
        st(ws,3,c,name,bg=BLUE,fg=WHITE,bold=True,size=10)
        st(ws,4,c,val,bg=PALE,bold=True,size=12,fmt=PCT if '율' in name else FMT)
        ws.merge_cells(start_row=3,start_column=c,end_row=3,end_column=c+1)
        ws.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+1)
    st(ws,6,1,'분기/시즌 분석이 가능한 필드',bg=GREEN,fg=WHITE,bold=True)
    for i,v in enumerate(['기준일=해당월 말일','분기=1Q~4Q','시즌=SS(3~8월), FW(9~다음해 2월)','1~2월은 전년도 FW'],1):
        st(ws,6+i,1,v,align='left')
    for col in range(1,11): ws.column_dimensions[get_column_letter(col)].width=16
    ws.sheet_view.showGridLines=False
    # 월별요약
    ws=wb.create_sheet('월별요약')
    monthly=aggregate(rows,['년-월'])
    write_table(ws, monthly, ['년-월']+METRICS+['순이익율'])
    # chart
    if len(monthly)>1:
        chart=LineChart(); chart.title='월별 매출/순이익 추이'; chart.y_axis.title='금액'; chart.x_axis.title='년월'
        data=Reference(ws,min_col=2,max_col=8,min_row=1,max_row=ws.max_row)
        cats=Reference(ws,min_col=1,min_row=2,max_row=ws.max_row)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.height=8; chart.width=22
        ws.add_chart(chart,'J2')
    # 분기분석
    ws=wb.create_sheet('분기분석')
    q=aggregate(rows,['년','분기'])
    write_table(ws,q,['년','분기']+METRICS+['순이익율'])
    # 시즌분석
    ws=wb.create_sheet('시즌분석')
    s=aggregate(rows,['시즌명','시즌연도','시즌'])
    write_table(ws,s,['시즌명','시즌연도','시즌']+METRICS+['순이익율'])
    # 매장별월별추이
    ws=wb.create_sheet('매장별월별추이')
    sm=aggregate(rows,['년-월','매장코드','매장명'])
    write_table(ws,sm,['년-월','매장코드','매장명']+METRICS+['순이익율'])
    # Top Bottom latest
    ws=wb.create_sheet('최신월_TOP_BOTTOM')
    latest_store=aggregate(latest_rows,['매장코드','매장명'])
    top=sorted(latest_store,key=lambda x:x['순이익'],reverse=True)[:20]
    bottom=sorted(latest_store,key=lambda x:x['순이익'])[:20]
    header(ws,1,['구분','순위','매장코드','매장명','판매금액','총경비','순이익','순이익율'])
    r=2
    for label,data in [('TOP',top),('BOTTOM',bottom)]:
        for i,row in enumerate(data,1):
            vals=[label,i,row['매장코드'],row['매장명'],row['판매금액'],row['총경비'],row['순이익'],row['순이익율']]
            for c,v in enumerate(vals,1): st(ws,r,c,v,fmt=PCT if c==8 else (FMT if c>=5 and c<=7 else None),align='left' if c==4 else 'center')
            r+=1
    for c,w in enumerate([10,8,12,22,14,14,14,10],1): ws.column_dimensions[get_column_letter(c)].width=w
    format_sheet(ws)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path

def run(base_dir: Path = None, ym: str = None):
    base_dir=Path(base_dir or Path(__file__).parent)
    db_path=base_dir/'DB'/'월별손익DB.xlsx'
    out_path=base_dir/'output'/'연도별_분기별_시즌별_손익분석.xlsx'
    make_report(db_path,out_path)
    print(f'✅ 손익 분석 보고서 생성: {out_path}')
    # ym이 안 넘어와도(예: history_report.py 단독 실행) DB의 최신월로 자동 보완한다.
    # 그렇지 않으면 output/YYYY-MM/ 사본이 갱신되지 않고 옛날 내용으로 남는 문제가 있었다.
    month_ym = ym
    if not month_ym:
        try:
            rows = load_db(db_path)
            month_ym = sorted({r['년-월'] for r in rows})[-1] if rows else None
        except Exception:
            month_ym = None
    if month_ym:
        try:
            import shutil
            month_out = base_dir/'output'/month_ym/'연도별_분기별_시즌별_손익분석.xlsx'
            month_out.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(out_path, month_out)
            print(f'   ↳ 월별 사본: {month_out}')
        except Exception as e:
            print(f'   ⚠️ 월별 사본 저장 오류: {e}')
    return out_path

if __name__=='__main__':
    run(Path(__file__).parent)
