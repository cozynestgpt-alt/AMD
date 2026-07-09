"""
중간관리판매수수료이체내역.xlsx 생성 모듈
- 중간관리 24개 매장의 송금액 이체 내역
- 송금액 = 개별통지문의 ⑦송금액과 동일
- 성명/은행/계좌: 사원마스터 기준
- 순서: 사원마스터(results) 순서와 동일
"""
import math
from pathlib import Path
from collections import defaultdict
from datetime import date

from common import find_input_one

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable,"-m","pip","install","openpyxl","--quiet"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# 은행명 약칭 → 정식명
BANK_MAP = {
    "신한":"신한은행","국민":"국민은행","기업":"기업은행","우리":"우리은행",
    "하나":"하나은행","농협":"농협은행","카카오":"카카오뱅크","케이":"케이뱅크",
}
def _nb(b):
    b = str(b or "").strip()
    return BANK_MAP.get(b, b)


def make_transfer_report(ym: str, base_dir: Path,
                         results: list,
                         expense_rows: dict,
                         master: list) -> Path:
    INPUT  = base_dir / "input" / ym
    OUTPUT = base_dir / "output" / ym
    OUTPUT.mkdir(parents=True, exist_ok=True)

    # ── 보조 데이터 로드 (정산월 폴더 안에서만 찾음) ─────────────────
    def _find_in(patterns):
        return find_input_one(patterns, base=INPUT)

    try:
        from amd_report import (load_phone_fee, load_tax_deduction,
                                load_delivery_fee, load_arba_subtotals)
        from expense_report import STORE_CODE_MAP
        phone_path = _find_in(["매장전화요금내역*.xlsx", "*전화요금*.xlsx"])
        tax_paths  = [_find_in(["급여상여명세서*일용직*.xlsx"]),
                      _find_in(["급여상여명세서*매장직*.xlsx"])]
        deliv_path = _find_in(["*로젠*고객직배*.xlsx", "*고객직배*.xlsx"])
        arba_path  = OUTPUT / "아르바이트_정산서.xlsx"
        phone_data = load_phone_fee(phone_path, ym) if phone_path else {}
        tax_data   = load_tax_deduction(tax_paths)
        deliv_data = load_delivery_fee(deliv_path) if deliv_path else {}
        arba_data  = load_arba_subtotals(arba_path)
    except Exception as e:
        print(f"     ⚠️  보조데이터 로드 오류: {e}")
        phone_data = {}; tax_data = {}; deliv_data = {}; arba_data = {}
        from expense_report import STORE_CODE_MAP

    mgr_shops   = {e["shop"] for e in master if e.get("pay_type","") == "중간관리"}
    event_shops = {e["shop"] for e in master if (e.get("note","") or "").strip() == "행사매장"}

    # 매장별 매니저 정보
    mgr_info = {}
    for e in master:
        if e.get("pay_type","")=="중간관리" and e.get("grade","")=="매니저" and e.get("name"):
            mgr_info[e["shop"]] = {
                "name":    e["name"],
                "bank":    _nb(e["bank"]),
                "account": e["account"] or "",
            }

    # 중간관리 매장 순서 (results 순서 유지)
    mgr_order = []; seen = set()
    for r in results:
        shop = r["emp"]["shop"]
        if r["emp"].get("pay_type","") == "중간관리" and shop not in seen:
            seen.add(shop); mgr_order.append(shop)

    # 매장별 수수료 집계
    shop_fee = defaultdict(int)
    shop_income = {}
    for r in results:
        emp, res = r["emp"], r["res"]
        if not emp["name"]: continue
        shop_fee[emp["shop"]] += emp["base_fee"] + res["commission"]
        if emp["shop"] not in shop_income:
            shop_income[emp["shop"]] = emp["income"]

    # ── 매장별 송금액 계산 (개별통지문과 동일 로직) ─────────────
    data_rows = []
    for shop in mgr_order:
        info = mgr_info.get(shop, {})
        er = expense_rows.get(shop, {})
        fee_c = math.ceil(shop_fee[shop] / 10) * 10
        공제 = (er.get("shortfall",0) + er.get("gift",0) + er.get("pos",0)
                - er.get("t_refund",0) + er.get("loss",0))
        공급 = fee_c + er.get("j_direct",0) - 공제
        inc  = shop_income.get(shop, "")
        vat  = int(공급 * 0.1) if inc == "사업소득" else 0
        지급 = 공급 + vat
        tel  = phone_data.get(shop, 0)
        deli = deliv_data.get(shop, 0)
        tax  = tax_data.get(shop, 0) if shop in STORE_CODE_MAP else 0
        기타 = tel + deli + tax
        arba = arba_data.get(shop, 0)
        송금 = 지급 - 기타 + arba
        data_rows.append({
            "shop": shop,
            "name": info.get("name", ""),
            "bank": info.get("bank", ""),
            "account": info.get("account", ""),
            "송금": 송금,
        })

    # ── 엑셀 새 파일로 생성 ─────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "이체내역서"
    ws.sheet_view.showGridLines = False

    THIN = Side(style="thin"); MED = Side(style="medium")
    FMT  = "#,##0"
    GRAY = "D9D9D9"

    def bdr(r1, c1, r2, c2):
        for r in range(r1, r2+1):
            for c in range(c1, c2+1):
                ws.cell(r,c).border = Border(
                    left   = MED if c==c1 else THIN,
                    right  = MED if c==c2 else THIN,
                    top    = MED if r==r1 else THIN,
                    bottom = MED if r==r2 else THIN)

    def st(r, c, v=None, bg=None, fg="000000", bold=False,
           size=11, align="center", fmt=None, font_name="맑은 고딕"):
        cell = ws.cell(r, c)
        if v is not None: cell.value = v
        cell.font      = Font(name=font_name, bold=bold, color=fg, size=size)
        cell.alignment = Alignment(horizontal=align, vertical="center")
        if bg: cell.fill = PatternFill("solid", fgColor=bg)
        if fmt: cell.number_format = fmt
        return cell

    # 열너비
    for col, w in {"A":5.1,"B":16.8,"C":9.0,"D":10.1,"E":19.5,"F":15.7,"G":15.8}.items():
        ws.column_dimensions[col].width = w

    # 행높이
    ws.row_dimensions[1].height = 35.1
    ws.row_dimensions[2].height = 19.2
    ws.row_dimensions[3].height = 6.0
    ws.row_dimensions[4].height = 17.1
    n = len(data_rows)
    for i in range(n): ws.row_dimensions[i+5].height = 17.1

    # 1행: 제목
    ws.merge_cells("A1:G1")
    st(1, 1, "판  매  수  수  료  이  체  내  역  서",
       bold=True, size=18, align="center", font_name="Microsoft NeoGothic")

    # 2행: 이체일자
    year, month = ym.split("-")
    이체일 = date(int(year), int(month)+1, 10)
    st(2, 6, "이체일자 :  ", align="right", size=11)
    ws.cell(2,7).value          = 이체일
    ws.cell(2,7).number_format  = "YYYY년 M월 D일"
    ws.cell(2,7).font           = Font(name="맑은 고딕", size=11)
    ws.cell(2,7).alignment      = Alignment(horizontal="center", vertical="center")

    # 4행: 헤더
    for c, h in [(1,"순번"),(2,"점포명"),(3,"예금주"),
                 (4,"은행명"),(5,"계좌번호"),(6,"금액"),(7,"비고")]:
        st(4, c, h, bg=GRAY, bold=True, size=11, align="center")
    bdr(4, 1, 4, 7)

    # 데이터
    for i, row in enumerate(data_rows):
        r = i + 5
        st(r, 1, i+1,            align="center", size=11)
        st(r, 2, row["shop"],     align="left",   size=11)
        st(r, 3, row["name"],     align="center", size=11)
        st(r, 4, row["bank"],     align="center", size=11)
        st(r, 5, row["account"],  align="center", size=11)
        st(r, 6, row["송금"],     align="right",  size=11, fmt=FMT)
        st(r, 7, "",              align="center", size=11)
    bdr(5, 1, n+4, 7)

    # 합계행
    sum_r = n + 5
    ws.row_dimensions[sum_r].height = 17.1
    total = sum(r["송금"] for r in data_rows)
    st(sum_r, 5, "합      계 :", align="right",  bold=True, size=11)
    st(sum_r, 6, total,           align="right",  bold=True, size=11, fmt=FMT)
    bdr(sum_r, 1, sum_r, 7)

    # 은행별 합계
    bank_sum = defaultdict(int); bank_cnt = defaultdict(int)
    for row in data_rows:
        if row["bank"]:
            bank_sum[row["bank"]] += row["송금"]
            bank_cnt[row["bank"]] += 1
    bank_order = [b for b in ["국민은행","신한은행","기업은행","우리은행","하나은행"]
                  if b in bank_sum]

    sec_r = sum_r + 2
    ws.row_dimensions[sec_r].height = 19.8
    ws.merge_cells(f"B{sec_r}:F{sec_r}")
    st(sec_r, 2, "2. 은행별 합계", bold=True, size=11, align="left")

    for bi, bk in enumerate(bank_order):
        br = sec_r + 1 + bi
        ws.row_dimensions[br].height = 20.4
        ws.merge_cells(f"B{br}:E{br}")
        st(br, 2, bk,              align="left",   size=11)
        st(br, 6, bank_cnt[bk],    align="center", size=11)
        st(br, 7, bank_sum[bk],    align="right",  bold=True, size=11, fmt=FMT)
        bdr(br, 2, br, 7)

    tot_r = sec_r + 1 + len(bank_order)
    ws.row_dimensions[tot_r].height = 20.4
    ws.merge_cells(f"B{tot_r}:E{tot_r}")
    st(tot_r, 2, "합계",                    align="center", bold=True, size=11)
    st(tot_r, 6, sum(bank_cnt.values()),    align="center", bold=True, size=11)
    st(tot_r, 7, sum(bank_sum.values()),    align="right",  bold=True, size=11, fmt=FMT)
    bdr(tot_r, 2, tot_r, 7)

    out_path = OUTPUT / "중간관리판매수수료이체내역.xlsx"
    wb.save(out_path)
    return out_path


def run(ym: str, base_dir: Path, results: list,
        expense_rows: dict, master: list):
    print(f"  💳 중간관리판매수수료이체내역 처리 중...")
    try:
        path = make_transfer_report(ym, base_dir, results, expense_rows, master)
        n = sum(1 for e in master if e.get("pay_type","") == "중간관리"
                and e.get("grade","") == "매니저" and e.get("name"))
        print(f"     ✅ 중간관리판매수수료이체내역.xlsx  ({n}개 매장)")
        return path
    except Exception as e:
        print(f"     ⚠️  이체내역 처리 오류: {e}")
        import traceback; traceback.print_exc()
        return None
