"""
매입세금계산서 자동 출력 모듈 v3
- 템플릿 서식/레이아웃 100% 유지
- #REF! / 수식 셀만 Python 계산값으로 교체
- 상단(공급받는자보관용, 행11~30) / 하단(공급자보관용, 행33~52) 동일 구조
"""
import math, copy, calendar
from pathlib import Path
from collections import defaultdict
from datetime import date

try:
    from openpyxl import load_workbook
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"])
    from openpyxl import load_workbook
    import openpyxl


# ── 공급자 기재사항 파일 로드 ────────────────────────────────
def load_supplier_info(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    본사 = {
        "상호명":    ws.cell(3, 4).value or "㈜더코지네스트컴퍼니",
        "사업자번호": ws.cell(3, 5).value or "211-81-83236",
        "대표자":    ws.cell(3, 6).value or "김창수",
        "주소":      ws.cell(3, 7).value or "서울시 강남구 논현2동 59-8",
        "업태":      ws.cell(3, 8).value or "제조,도소매",
        "종목":      ws.cell(3, 9).value or "침구 및 관련제품",
    }
    data = {}
    for r in range(4, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if not isinstance(no, (int, float)):
            continue
        전산명 = str(ws.cell(r, 3).value or "").strip()
        if not 전산명:
            continue
        data[전산명] = {
            "no":        int(no),
            "상호명":    ws.cell(r, 4).value or "",
            "사업자번호": ws.cell(r, 5).value or "",
            "대표자":    ws.cell(r, 6).value or "",
            "주소":      ws.cell(r, 7).value or "",
            "업태":      ws.cell(r, 8).value or "",
            "종목":      ws.cell(r, 9).value or "",
        }
    return 본사, data


def _sv(ws, r, c, v):
    """MergedCell 안전 셀 입력"""
    try:
        cell = ws.cell(r, c)
        if cell.__class__.__name__ != "MergedCell":
            cell.value = v
    except Exception:
        pass


def _digits(ws, number, row, col_start, col_end):
    """
    숫자를 자릿수별로 오른쪽 정렬해서 셀에 입력
    col_end(일의 자리) 기준으로 오른쪽 → 왼쪽으로 채움
    """
    s = str(int(number))
    for i, d in enumerate(reversed(s)):
        col = col_end - i
        if col_start <= col <= col_end:
            _sv(ws, row, col, d)


def _fill_block(ws, br, info, 공급가액, 세액, 작성일):
    """
    세금계산서 1블록 값 입력
    br=11 → 상단(공급받는자보관용)
    br=33 → 하단(공급자보관용)

    템플릿 레이아웃 기준 (br=11 상단):
      행14: 공급자 등록번호(F14), 공급받는자 등록번호(V14~AG14)
      행15: 공급자 상호(F15), 성명(N15)
      행17: 공급자 사업장주소(F17)
      행19: 공급자 업태(F19), 종목(L19)
      행22: 작성년(B22), 월(D22), 일(E22), 공급가액자릿수(G22~Q22), 세액자릿수(R22~AA22)
      행24: 월(B24), 일(C24), 품목=판매수수료, 공급가액(T24), 세액(Z24)
      행29: 합계금액(B29)
    하단(br=33)은 동일 오프셋
    """
    # ── 공급자 정보 ──────────────────────────────────────────
    _sv(ws, br + 3, 6,  info.get("사업자번호", ""))   # F14/F36
    _sv(ws, br + 4, 6,  info.get("상호명", ""))        # F15/F37
    _sv(ws, br + 4, 14, info.get("대표자", ""))        # N15/N37
    _sv(ws, br + 6, 6,  info.get("주소", ""))          # F17/F39
    _sv(ws, br + 8, 6,  info.get("업태", ""))          # F19/F41
    _sv(ws, br + 8, 12, info.get("종목", ""))          # L19/L41

    # ── 작성일 / 공급가액·세액 자릿수 ────────────────────────
    _sv(ws, br + 11, 2, 작성일.year)    # B22/B44
    _sv(ws, br + 11, 4, 작성일.month)   # D22/D44
    _sv(ws, br + 11, 5, 작성일.day)     # E22/E44

    # 공급가액 자릿수: G22~Q22 (열7~17)
    _digits(ws, 공급가액, br + 11, 7, 17)
    # 세액 자릿수: R22~AA22 (열18~27)
    _digits(ws, 세액, br + 11, 18, 27)

    # ── 품목행 ────────────────────────────────────────────────
    _sv(ws, br + 13, 2, 작성일.month)   # B24/B46 월
    _sv(ws, br + 13, 3, 작성일.day)     # C24/C46 일
    # D24/D46 "판매수수료" 는 템플릿에 이미 있음
    _sv(ws, br + 13, 20, 공급가액)      # T24/T46 공급가액
    ws.cell(br + 13, 20).number_format = "#,##0"
    _sv(ws, br + 13, 26, 세액)          # Z24/Z46 세액
    ws.cell(br + 13, 26).number_format = "#,##0"

    # ── 합계금액 ─────────────────────────────────────────────
    _sv(ws, br + 18, 2, 공급가액 + 세액)  # B29/B51
    ws.cell(br + 18, 2).number_format = "#,##0"


def _clear_ref(ws):
    """
    #REF! 값(수식 결과)과 수식 문자열 모두 None으로 초기화
    data_only=True 로 열면 이미 평가된 값이 들어있음
    """
    for row in ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            v = cell.value
            if v is None:
                continue
            # 수식 문자열 (data_only=False 로 열었을 때)
            if isinstance(v, str) and v.startswith("="):
                cell.value = None
            # #REF! 결과값
            elif isinstance(v, str) and "#REF!" in v:
                cell.value = None


def _apply_page_setup(ws, print_area="A10:AH53"):
    ws.print_area = print_area
    ws.page_setup.paperSize   = 9
    ws.page_setup.orientation = "portrait"
    ws.page_setup.scale       = 99
    # fitToPage 미설정 — 템플릿 원본과 동일하게 scale=99 고정
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered   = True
    ws.page_margins.left   = 0.19685039370078741
    ws.page_margins.right  = 0.19685039370078741
    ws.page_margins.top    = 0.59055118110236227
    ws.page_margins.bottom = 0.59055118110236227
    ws.page_margins.header = 0.31496062992125984
    ws.page_margins.footer = 0.31496062992125984


def make_tax_invoice(ym: str, base_dir: Path,
                     results: list, expense_rows: dict, master: list) -> Path:
    INPUT    = base_dir / "input"
    TMPL_DIR = base_dir / "templates"
    OUTPUT   = base_dir / "output" / ym
    OUTPUT.mkdir(parents=True, exist_ok=True)

    info_path = next(iter(INPUT.glob("중간관리매장_매입세금계산서*.xlsx")), None)
    tmpl_path = TMPL_DIR / "세금계산서출력_템플릿.xlsx"

    if not info_path:
        print(f"     ⚠️  기재사항정보 파일 없음: input/중간관리매장_매입세금계산서_기재사항정보.xlsx")
        return None
    if not tmpl_path.exists():
        print(f"     ⚠️  템플릿 없음: templates/세금계산서출력_템플릿.xlsx")
        return None

    본사_info, supplier_info = load_supplier_info(info_path)

    try:
        from common import STORE_CODE_MAP
    except ImportError:
        from expense_report import STORE_CODE_MAP

    event_shops = {e["shop"] for e in master if (e.get("note", "") or "").strip() == "행사매장"}

    # 매장별 수수료 합산
    shop_fee    = defaultdict(int)
    shop_income = {}
    for r in results:
        emp, res = r["emp"], r["res"]
        if not emp["name"]:
            continue
        shop_fee[emp["shop"]] += emp["base_fee"] + res["commission"]
        if emp["shop"] not in shop_income:
            shop_income[emp["shop"]] = emp["income"]

    # 중간관리 + 사업소득 매장 순서 유지
    seen = set()
    invoice_shops = []
    for r in results:
        shop = r["emp"]["shop"]
        if (r["emp"].get("pay_type", "") == "중간관리"
                and shop_income.get(shop) == "사업소득"
                and shop not in seen):
            seen.add(shop)
            invoice_shops.append(shop)

    year, month = int(ym[:4]), int(ym[5:])
    작성일 = date(year, month, calendar.monthrange(year, month)[1])

    # ── 출력 워크북 생성 ─────────────────────────────────────
    wb_out = openpyxl.Workbook()
    del wb_out["Sheet"]

    shop_values = {}

    for shop in invoice_shops:
        info = supplier_info.get(shop, {})
        er   = expense_rows.get(shop, {})

        fee_c    = math.ceil(shop_fee[shop] / 10) * 10
        공제     = (er.get("shortfall", 0) + er.get("gift", 0) + er.get("pos", 0)
                    - er.get("t_refund", 0) + er.get("loss", 0))
        공급가액 = fee_c + er.get("j_direct", 0) - 공제
        세액     = int(공급가액 * 0.1)
        shop_values[shop] = {"공급가액": 공급가액, "세액": 세액}

        # ── 템플릿을 data_only=False 로 열어 서식 복사, 수식 제거 ──
        wb_tmpl = load_workbook(tmpl_path, data_only=False)
        ws_tmpl = wb_tmpl["세금계산서(출력)"]
        _clear_ref(ws_tmpl)   # #REF! / 수식 모두 None 처리

        # 값 주입
        _fill_block(ws_tmpl, 11, info, 공급가액, 세액, 작성일)   # 상단
        _fill_block(ws_tmpl, 33, info, 공급가액, 세액, 작성일)   # 하단

        # ── 출력 워크북에 시트 복사 ──────────────────────────
        순번    = info.get("no", invoice_shops.index(shop) + 1)
        sheet_nm = f"{순번:02d}_{shop}"[:31]
        ws_out   = wb_out.create_sheet(sheet_nm)

        # 열너비: 템플릿과 동일하게 A~AH(1~34열) 전체 2.59765625
        from openpyxl.utils import get_column_letter
        for i in range(1, 35):
            ws_out.column_dimensions[get_column_letter(i)].width = 2.59765625
        # 행높이
        for row, rd in ws_tmpl.row_dimensions.items():
            ws_out.row_dimensions[row].height = rd.height

        # 셀 스타일+값 복사
        for sr in ws_tmpl.iter_rows():
            for sc in sr:
                if sc.__class__.__name__ == "MergedCell":
                    continue
                dc = ws_out.cell(sc.row, sc.column)
                dc.value = sc.value
                if sc.has_style:
                    dc.font          = copy.copy(sc.font)
                    dc.fill          = copy.copy(sc.fill)
                    dc.border        = copy.copy(sc.border)
                    dc.alignment     = copy.copy(sc.alignment)
                    dc.number_format = sc.number_format

        # 병합셀 (스타일 복사 후)
        for mr in ws_tmpl.merged_cells.ranges:
            try:
                ws_out.merge_cells(str(mr))
            except Exception:
                pass

        _apply_page_setup(ws_out)
        ws_out.sheet_view.showGridLines = False

    # ── 매입세금계산서_list 시트 복사 + 값 입력 ──────────────
    wb_tmpl2 = load_workbook(tmpl_path, data_only=False)
    ws_src   = wb_tmpl2["매입세금계산서_list"]
    ws_list  = wb_out.create_sheet("매입세금계산서_list")

    for col, cd in ws_src.column_dimensions.items():
        ws_list.column_dimensions[col].width = cd.width
    for row, rd in ws_src.row_dimensions.items():
        ws_list.row_dimensions[row].height = rd.height
    for sr in ws_src.iter_rows():
        for sc in sr:
            if sc.__class__.__name__ == "MergedCell":
                continue
            dc = ws_list.cell(sc.row, sc.column)
            v = sc.value
            # 수식·#REF! 제거
            if isinstance(v, str) and (v.startswith("=") or "#REF!" in v):
                dc.value = None
            else:
                dc.value = v
            if sc.has_style:
                dc.font          = copy.copy(sc.font)
                dc.fill          = copy.copy(sc.fill)
                dc.border        = copy.copy(sc.border)
                dc.alignment     = copy.copy(sc.alignment)
                dc.number_format = sc.number_format
    for mr in ws_src.merged_cells.ranges:
        try:
            ws_list.merge_cells(str(mr))
        except Exception:
            pass

    FMT = "#,##0"
    total_공급 = 0
    total_세액 = 0
    for i, shop in enumerate(invoice_shops):
        row      = 6 + i
        공급가액 = shop_values[shop]["공급가액"]
        세액     = shop_values[shop]["세액"]
        합계     = 공급가액 + 세액
        for c, v in [(12, 공급가액), (13, 공급가액), (14, 세액), (15, 합계)]:
            try:
                ws_list.cell(row, c).value        = v
                ws_list.cell(row, c).number_format = FMT
            except Exception:
                pass
        total_공급 += 공급가액
        total_세액 += 세액

    for c, v in [(12, total_공급), (13, total_공급),
                 (14, total_세액), (15, total_공급 + total_세액)]:
        try:
            ws_list.cell(2, c).value        = v
            ws_list.cell(2, c).number_format = FMT
        except Exception:
            pass
    ws_list.sheet_view.showGridLines = False

    out_path = OUTPUT / f"세금계산서_{ym}.xlsx"
    wb_out.save(out_path)
    return out_path


def run(ym, base_dir, results, expense_rows, master):
    print(f"  🧾 세금계산서 처리 중...")
    try:
        path = make_tax_invoice(ym, base_dir, results, expense_rows, master)
        if path:
            from collections import defaultdict
            shop_fee    = defaultdict(int)
            shop_income = {}
            for r in results:
                emp = r["emp"]
                if not emp["name"]: continue
                shop_fee[emp["shop"]] += emp["base_fee"] + r["res"]["commission"]
                if emp["shop"] not in shop_income:
                    shop_income[emp["shop"]] = emp["income"]
            cnt = sum(1 for s, inc in shop_income.items()
                      if inc == "사업소득"
                      and any(r["emp"]["shop"] == s and r["emp"].get("pay_type") == "중간관리"
                              for r in results))
            print(f"     ✅ 세금계산서_{ym}.xlsx  ({cnt}개 매장)")
        return path
    except Exception as e:
        print(f"     ⚠️  세금계산서 처리 오류: {e}")
        import traceback; traceback.print_exc()
        return None
