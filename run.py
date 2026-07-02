"""
판매수수료 자동계산 시스템
사용법: python run.py
실행 후 정산 월(YYYY-MM) 입력 → 자동 처리
"""

import os, sys, subprocess, shutil
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("필요한 라이브러리를 설치합니다...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# ── 경로 설정 ──────────────────────────────────────────────
BASE   = Path(__file__).parent
INPUT  = BASE / "input"
OUTPUT = BASE / "output"
TMPL   = BASE / "templates"

# ── 색상 팔레트 ────────────────────────────────────────────
NAVY        = "1F3864"
BLUE        = "2E5EAA"
LIGHT_BLUE  = "D9E2F3"
PALE_BLUE   = "EEF3FB"
GREEN       = "1E6B3C"
LIGHT_GREEN = "E2EFDA"
AMBER       = "FFF2CC"
AMBER_DARK  = "BF8F00"
WHITE       = "FFFFFF"
GRAY        = "F2F2F2"
GRAY_MID    = "D6D6D6"
RED_LIGHT   = "FFE0E0"
CORAL       = "C00000"
INPUT_BLUE  = "0000FF"
LINK_GREEN  = "008000"
AUTO_BLACK  = "000000"

FMT_COMMA = '#,##0'
FMT_PCT   = '0.0%'

# ══════════════════════════════════════════════════════════
# 공통 스타일 함수
# ══════════════════════════════════════════════════════════
def st(ws, row, col, val=None, bg=None, fg="000000", bold=False,
       size=9, align="center", wrap=False, fmt=None, italic=False):
    c = ws.cell(row=row, column=col)
    if val is not None:
        c.value = val
    c.font      = Font(name="맑은 고딕", bold=bold, color=fg, size=size, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if fmt:
        c.number_format = fmt
    return c

def borders(ws, r1, c1, r2, c2, outer="medium", inner="thin"):
    tk = Side(style=outer); tn = Side(style=inner)
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            ws.cell(r, c).border = Border(
                left  =tk if c==c1 else tn,
                right =tk if c==c2 else tn,
                top   =tk if r==r1 else tn,
                bottom=tk if r==r2 else tn,
            )

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ══════════════════════════════════════════════════════════
# 데이터 읽기
# ══════════════════════════════════════════════════════════
def load_master(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    employees = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:                          # NO 없으면 종료
            break
        if not isinstance(row[0], (int,float)): # NO가 숫자가 아니면(합계행 등) 건너뜀
            continue
        # 성명 없어도 비고='행사매장'이면 유지 (POS환급 제외 등 행사매장 판별 필요)
        if not row[2] and str(row[11] or "").strip() != "행사매장":
            continue
        employees.append({
            "no":       row[0],
            "shop":     row[1] or "",
            "name":     row[2] or "",
            "grade":    row[3] or "",
            "income":   row[4] or "",
            "pay_type": row[5] or "",
            "bank":     row[6] or "",
            "account":  row[7] or "",
            "base_fee": row[8] or 0,
            "rate_off": row[9] or 0,
            "rate_on":  row[10] or 0,
            "note":     row[11] or "",
        })
    return employees

def load_deduction(path: Path) -> dict:
    """
    ◈매장공제건 집계 파일 → {매장명: {shortfall, loss, gift, pos}}
    시트명: '매장별공제금액(2)' 또는 '매장별공제금액 (2)'
    컬럼: C=매장명, D=덜받음, E=재고LOSS, F=임의사은품, G=임의상품권대체할인,
          H=임의사은품지급합계, I=임의추가할인, J=POS오류정정, K=POS수혜매장
    """
    if not path or not path.exists():
        return {}
    try:
        import pandas as pd
        # 시트명 유연하게 탐지
        xl = pd.ExcelFile(path)
        sheet = next((s for s in xl.sheet_names
                      if "매장별공제금액" in s), xl.sheet_names[0])
        df = pd.read_excel(path, sheet_name=sheet, header=0)
        data = {}
        for _, row in df.iterrows():
            shop = str(row.iloc[2]).strip() if not pd.isna(row.iloc[2]) else ""
            if not shop or shop in ("매장명","nan","합계"): continue
            def v(col):
                val = row.iloc[col]
                return int(val) if not pd.isna(val) and val else 0
            # F(idx5)+G(idx6) = 임의사은품 합산
            gift = v(5) + v(6)
            data[shop] = {
                "shortfall": v(3),   # D: 덜받음
                "loss":      v(4),   # E: 재고LOSS
                "gift":      gift,   # F+G: 임의사은품+상품권대체할인
                "pos":       v(9),   # J: POS오류정정
            }
        return data
    except Exception as e:
        print(f"  ⚠️  공제건집계 파일 읽기 오류: {e}")
        return {}

# ══════════════════════════════════════════════════════════
# 수수료 계산
# ══════════════════════════════════════════════════════════
def calc(emp: dict, sales: dict, ded: dict = None) -> dict:
    """
    ded : load_deduction()의 결과 {shop: {shortfall, loss, gift, pos}}
          → 매장공제건집계 파일에서 공제항목 읽음
    매출 데이터는 sales (영*판매*.xlsx) 에서만 읽음
    """
    s = sales.get(emp["shop"], {})
    d = (ded or {}).get(emp["shop"], {})   # 매장공제건집계

    off_sales   = s.get("off", 0)
    on_sales    = s.get("on",  0)
    total_sales = off_sales + on_sales

    base_fee  = emp["base_fee"]
    rate_off  = emp["rate_off"]
    rate_on   = emp["rate_on"]

    # 수수료 계산: 오프+온 합산 후 10원 미만 올림
    if rate_off > 0:
        import math as _math
        commission = _math.ceil(
            (off_sales * rate_off + on_sales * rate_on) / 10
        ) * 10
    else:
        commission = 0

    arba      = 0       # 아르바이트: 아르바이트_정산서에서 별도 처리
    direct    = d.get("direct",    0)   # 직배비: expense_report에서
    shortfall = d.get("shortfall", 0)   # 덜받음
    loss      = d.get("loss",      0)   # 유통하자LOSS
    gift      = d.get("gift",      0)   # 임의사은품
    pos       = d.get("pos",       0)   # POS공제
    deduct    = shortfall + loss + gift + pos

    subtotal  = base_fee + commission + arba + direct - deduct

    # 부가세 (사업소득자만)
    vat = int(subtotal * 0.1) if emp["income"] == "사업소득" else 0
    total_pay = subtotal + vat

    return {
        "off_sales":   off_sales,
        "on_sales":    on_sales,
        "total_sales": total_sales,
        "base_fee":    base_fee,
        "commission":  commission,
        "arba":        arba,
        "direct":      direct,
        "shortfall":   shortfall,
        "loss":        loss,
        "gift":        gift,
        "pos":         pos,
        "deduct":      deduct,
        "subtotal":    subtotal,
        "vat":         vat,
        "total_pay":   total_pay,
        "note":        "",
    }

# ══════════════════════════════════════════════════════════
# ══════════════════════════════════════════════════════════
# 출력 3: 개별통지문.xlsx  (사원별 시트 + 전체 인쇄 설정)
# ══════════════════════════════════════════════════════════
def make_notice(results: list, ym: str, out_dir: Path,
                sales_detail: dict = None,
                expense_rows: dict = None):
    """
    사업소득자만 개별 통지문 생성 — 템플릿(개별통지문_템플릿.xlsx) 기반
    새 템플릿 구조 (G열=값입력):
      (2,2) 제목,  (4,5) 지급일자,  (5,5) 점포명
      (7,7) 정상매출,  (8,7) 행사매출
      (11,7) 매니저수수료,  (12,7) 지원금(직배비·경비)
      (14,7) 덜받음,  (15,7) 사은품,  (16,7) POS공제,  (17,7) POS환급,  (18,7) LOSS
      (23,7) 전화요금,  (24,7) 직배비공제(V+),  (25,7) 기타공제
    """
    import math, copy
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    biz_results = [r for r in results if r["emp"]["income"] == "사업소득"]

    TMPL_PATH = BASE / "templates" / "개별통지문_템플릿.xlsx"
    if not TMPL_PATH.exists():
        print(f"     ⚠️  템플릿 없음: {TMPL_PATH}")
        return None

    year, month = ym.split("-")
    ORANGE_BG = PatternFill("solid", fgColor="F4B460")
    FMT = "#,##0"

    # 전화요금·세액공제·직배비공제 데이터 로드
    try:
        from amd_report import (load_phone_fee, load_tax_deduction,
                                load_delivery_fee, DELIVERY_FEE_EXCLUDE)
        _INPUT = BASE / "input"
        _phone_path  = next(iter(_INPUT.glob("매장전화요금내역.xlsx")), None) or                        next(iter(_INPUT.glob("*전화요금*.xlsx")), None)
        _tax_paths   = [
            next(iter(_INPUT.glob("급여상여명세서_일용직*.xlsx")), None),
            next(iter(_INPUT.glob("급여상여명세서_매장직*.xlsx")), None),
        ]
        _deliv_path  = next(iter(_INPUT.glob("*로젠*고객직배*.xlsx")), None) or                        next(iter(_INPUT.glob("*고객직배*.xlsx")), None)
        _phone_data  = load_phone_fee(_phone_path, ym) if _phone_path else {}
        _tax_data    = load_tax_deduction(_tax_paths)
        _deliv_data2 = load_delivery_fee(_deliv_path) if _deliv_path else {}
    except Exception:
        _phone_data = {}; _tax_data = {}; _deliv_data2 = {}

    _mgr_shops = {r["emp"]["shop"] for r in results
                  if r["emp"].get("pay_type","") == "중간관리"}

    wb = load_workbook(TMPL_PATH)
    del wb[wb.sheetnames[0]]   # 기본 시트 삭제 후 재생성

    idx = 0
    for r in biz_results:
        emp, res = r["emp"], r["res"]
        shop = emp["shop"]; name = emp["name"]

        # 매출
        sd   = (sales_detail or {}).get(shop, {})
        off_n= sd.get("off_normal",0); off_e=sd.get("off_event",0)
        on_n = sd.get("on_normal",0);  on_e =sd.get("on_event",0)
        정상매출 = off_n + on_n
        행사매출 = off_e + on_e

        # 수수료 (10원 미만 올림)
        mgr_fee = math.ceil((res["commission"] + res["base_fee"]) / 10) * 10

        # 경비
        er       = (expense_rows or {}).get(shop, {})
        j_direct = er.get("j_direct",  0)
        shortfall= er.get("shortfall", 0)
        gift     = er.get("gift",      0)
        pos_ded  = er.get("pos",       0)
        pos_ref  = er.get("t_refund",  0)
        loss     = er.get("loss",      0)

        # 기타공제
        tel_fee    = _phone_data.get(shop, 0) if shop in _mgr_shops else 0
        direct_fee = _deliv_data2.get(shop, 0) if shop in _mgr_shops else 0
        tax_fee    = _tax_data.get(shop, 0)
        기타공제합계= tel_fee + direct_fee + tax_fee

        # 계산 (수식 없이 값으로)
        지급합계  = mgr_fee + j_direct
        공제합계  = shortfall + gift + pos_ded - pos_ref + loss
        공급가액  = 지급합계 - 공제합계
        부가세    = int(공급가액 * 0.1)
        지급할총액 = 공급가액 + 부가세
        송금액    = 지급할총액 - 기타공제합계

        # ── 시트 생성 (스타일 먼저, 병합 나중) ──────────────────
        sheet_nm  = f"{idx+1:02d}_{name}"[:31]
        wb_tmp    = load_workbook(TMPL_PATH)
        ws_src    = wb_tmp.active
        ws        = wb.create_sheet(sheet_nm)

        # 열너비 / 행높이
        for col, cd in ws_src.column_dimensions.items():
            ws.column_dimensions[col].width = cd.width
        for row, rd in ws_src.row_dimensions.items():
            ws.row_dimensions[row].height = rd.height

        # 셀 스타일 복사 (MergedCell 건너뜀)
        for sr in ws_src.iter_rows():
            for sc in sr:
                if sc.__class__.__name__ == "MergedCell":
                    continue
                dc = ws.cell(sc.row, sc.column)
                dc.value = 0 if isinstance(sc.value,str) and sc.value.startswith("=") else sc.value
                if sc.has_style:
                    dc.font      = copy.copy(sc.font)
                    dc.fill      = copy.copy(sc.fill)
                    dc.border    = copy.copy(sc.border)
                    dc.alignment = copy.copy(sc.alignment)
                    dc.number_format = sc.number_format

        # 값 주입 — 병합 전에 (G열=7열 기준)
        ws.cell(2,2).value = f"{year}년 {int(month)}월 판매수수료 내역서"
        ws.cell(4,5).value = f"{year}년  {int(month)+1}월  10일"
        ws.cell(5,5).value = shop

        def sv(r, v):
            c = ws.cell(r, 7)
            c.value = v if v else 0
            c.number_format = FMT
            if isinstance(v,(int,float)) and v < 0:
                c.font = Font(name=c.font.name or "맑은 고딕",
                              bold=c.font.bold, size=c.font.size or 11,
                              color="C00000")

        sv(7,  정상매출)
        sv(8,  행사매출)
        sv(9,  정상매출 + 행사매출)
        sv(11, mgr_fee)
        sv(12, j_direct)
        sv(13, 지급합계)
        sv(14, shortfall)
        sv(15, gift)
        sv(16, pos_ded)
        sv(17, pos_ref)
        sv(18, loss)
        sv(19, 공제합계)
        sv(20, 공급가액)
        sv(21, 부가세)
        ws.cell(22,7).value = 지급할총액
        ws.cell(22,7).number_format = FMT
        ws.cell(22,7).fill  = ORANGE_BG
        sv(23, tel_fee)
        sv(24, direct_fee)
        sv(25, tax_fee)
        sv(27, 기타공제합계)
        ws.cell(27,7).fill = ORANGE_BG
        ws.cell(28,7).value = 송금액
        ws.cell(28,7).number_format = FMT
        ws.cell(28,7).fill  = ORANGE_BG

        # 병합셀 (스타일+값 설정 후)
        for mr in ws_src.merged_cells.ranges:
            ws.merge_cells(str(mr))

        ws.sheet_view.showGridLines = False
        ws.page_setup.paperSize     = 9
        ws.print_area               = "A1:G32"
        ws.page_margins.left  = 0.4; ws.page_margins.right  = 0.4
        ws.page_margins.top   = 0.5; ws.page_margins.bottom = 0.5
        idx += 1

    # 목차 시트 (맨 앞)
    ws_idx = wb.create_sheet("📋목차")
    wb.move_sheet("📋목차", offset=-len(wb.sheetnames)+1)
    ws_idx.sheet_view.showGridLines = False
    from openpyxl.utils import get_column_letter
    for i,w in enumerate([5,28,12,16],1):
        ws_idx.column_dimensions[get_column_letter(i)].width = w
    from openpyxl.styles import Font as F2, PatternFill as P2, Alignment as A2
    NAVY="1F3864"; WHITE="FFFFFF"; GREEN="1E6B3C"; BLUE="2E5EAA"
    def si(r,c,v,bg=None,fg="000000",bold=False,size=9,align="center",fmt=None):
        cell=ws_idx.cell(r,c,v)
        cell.font=F2(name="맑은 고딕",bold=bold,color=fg,size=size)
        cell.alignment=A2(horizontal=align,vertical="center")
        if bg: cell.fill=P2("solid",fgColor=bg)
        if fmt: cell.number_format=fmt
    ws_idx.merge_cells("A1:D1")
    si(1,1,f"개별 통지문 목차  ▶  {ym}  (사업소득자 {len(biz_results)}명)",
       bg=NAVY,fg=WHITE,bold=True,size=12)
    ws_idx.row_dimensions[1].height=26
    for ci,h in enumerate(["NO","매장명·성명","지급할총액","시트이동"],1):
        si(2,ci,h,bg=BLUE,fg=WHITE,bold=True,size=9)
    ws_idx.row_dimensions[2].height=20
    GRAY="F2F2F2"
    for ii,r in enumerate(biz_results):
        emp,res=r["emp"],r["res"]
        sheet_nm=f"{ii+1:02d}_{emp['name']}"[:31]
        rb=WHITE if ii%2==0 else GRAY
        er=(expense_rows or {}).get(emp["shop"],{})
        mgr_f=math.ceil((res["commission"]+res["base_fee"])/10)*10
        j_d=er.get("j_direct",0)
        short=er.get("shortfall",0); gift_=er.get("gift",0)
        pos_d=er.get("pos",0); pos_r=er.get("t_refund",0); loss_=er.get("loss",0)
        공제=short+gift_+pos_d-pos_r+loss_
        공급=mgr_f+j_d-공제; vat=int(공급*0.1); total=공급+vat
        si(ii+3,1,ii+1,bg=rb,size=9)
        si(ii+3,2,f"{emp['shop']} {emp['name']}",bg=rb,size=9,align="left")
        si(ii+3,3,total,bg=rb,size=9,fmt="#,##0")
        cell=ws_idx.cell(ii+3,4,"→ 명세서")
        cell.hyperlink=f"#{sheet_nm}!A1"
        cell.font=F2(name="맑은 고딕",color="0000FF",underline="single",size=9)
        cell.alignment=A2(horizontal="center",vertical="center")
        ws_idx.row_dimensions[ii+3].height=17

    path = out_dir / "개별통지문.xlsx"
    wb.save(path)
    return path


# ══════════════════════════════════════════════════════════
# 출력 4: 처리결과_요약.txt
# ══════════════════════════════════════════════════════════
def make_summary(results: list, ym: str, out_dir: Path,
                 missing_sales: list,
                 expense_rows: dict = None, master: list = None):
    import math as _math
    from collections import defaultdict as _dd
    biz_cnt  = sum(1 for r in results if r["emp"]["income"]=="사업소득")
    work_cnt = sum(1 for r in results if r["emp"]["income"]!="사업소득")

    # ── 실무 기준 수치 계산 (판매수수료집계와 동일) ─────────────
    try:
        from expense_report import STORE_CODE_MAP as _SCM
        from amd_report import (load_phone_fee, load_tax_deduction,
                                load_delivery_fee, load_arba_subtotals)
        _ev = {e["shop"] for e in (master or [])
               if (e.get("note","") or "").strip()=="행사매장"}
        _mgr = {e["shop"] for e in (master or [])
                if e.get("pay_type","")=="중간관리"}
        _phone = load_phone_fee(
            next(iter(INPUT.glob("매장전화요금내역.xlsx")),None), ym)
        _tax = load_tax_deduction([
            next(iter(INPUT.glob("급여상여명세서_일용직*.xlsx")),None),
            next(iter(INPUT.glob("급여상여명세서_매장직*.xlsx")),None)])
        _deliv = load_delivery_fee(
            next(iter(INPUT.glob("*로젠*고객직배*.xlsx")),None))
        _arba  = load_arba_subtotals(out_dir/"아르바이트_정산서.xlsx")

        _shop_fee = _dd(int); _shop_inc = {}
        for _r in results:
            _e, _rs = _r["emp"], _r["res"]
            if not _e["name"]: continue
            _shop_fee[_e["shop"]] += _e["base_fee"] + _rs["commission"]
            if _e["shop"] not in _shop_inc:
                _shop_inc[_e["shop"]] = _e["income"]

        _총공급 = 0; _총부가세 = 0; _총기타공제 = 0
        for _shop, _inc in _shop_inc.items():
            _er = (expense_rows or {}).get(_shop, {})
            _fc = _math.ceil(_shop_fee[_shop]/10)*10
            _공제 = (_er.get("shortfall",0)+_er.get("gift",0)+_er.get("pos",0)
                     -_er.get("t_refund",0)+_er.get("loss",0))
            _공급 = _fc + _er.get("j_direct",0) - _공제
            _공급 += _arba.get(_shop, 0)   # 아르바이트 포함
            _vat  = int(_공급 * 0.1) if _inc=="사업소득" else 0
            _총공급  += _공급
            _총부가세 += _vat
            _tel  = _phone.get(_shop,0) if _shop in _mgr else 0
            _deli = _deliv.get(_shop,0) if _shop in _mgr else 0
            _tx   = _tax.get(_shop,0)   if _shop in _SCM else 0
            _총기타공제 += _tel + _deli + _tx

        # 행사매장 아르바이트 + 세액공제
        for _shop, _amt in _arba.items():
            if _shop in _ev:
                _총공급 += _amt
                _총기타공제 += _tax.get(_shop,0) if _shop in _SCM else 0

        _총지급액 = _총공급 + _총부가세
        _송금액   = _총지급액 - _총기타공제
        _total_vat = _총부가세

    except Exception as _err:
        # fallback: calc() 기반 (직배비·아르바이트 제외)
        _총지급액  = sum(r["res"]["total_pay"] for r in results)
        _total_vat = sum(r["res"]["vat"]       for r in results)
        _송금액    = _총지급액
        _총기타공제 = 0

    lines = [
        "="*60,
        f"  판매수수료 처리 결과 요약  ▶  {ym}",
        f"  실행일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "="*60,
        f"  처리 인원     : {len(results):>5}명",
        f"  사업소득자    : {biz_cnt:>5}명",
        f"  근로소득자    : {work_cnt:>5}명",
        "-"*60,
        f"  ⑦ 총 지급 합계: {_총지급액:>15,}원",
        f"     부가세 합계 : {_total_vat:>15,}원",
        f"  ⑨ 실 송 금 액 : {_송금액:>15,}원",
        "="*60,
    ]
    if missing_sales:
        lines += ["", "⚠️  매출 파일에 없는 매장 (매출 0으로 처리):"]
        for n in missing_sales:
            lines.append(f"    - {n}")
    lines += ["", "생성된 파일:", f"  📂 {out_dir}"]

    txt = "\n".join(lines)
    path = out_dir / "처리결과_요약.txt"
    path.write_text(txt, encoding="utf-8-sig")
    print(txt)
    return path

# ══════════════════════════════════════════════════════════
# PDF 안내 출력
# ══════════════════════════════════════════════════════════
def print_pdf_guide(notice_path: Path):
    print("\n" + "="*60)
    print("  📄 개별 통지문 PDF 저장 방법")
    print("="*60)
    print(f"\n  파일: {notice_path.name}")
    print("\n  【전체 PDF 저장】")
    print("    1. 개별통지문.xlsx 파일을 Excel로 엽니다")
    print("    2. 첫 번째 시트(📋목차) 탭을 클릭")
    print("    3. Shift 키를 누른 채 마지막 사원 시트 탭 클릭")
    print("       → 모든 시트가 선택됩니다")
    print("    4. 파일 → 내보내기 → PDF/XPS 문서 만들기")
    print("       또는  Ctrl+P → 프린터: 'Microsoft Print to PDF'")
    print("\n  【개별 사원 PDF 저장】")
    print("    1. 해당 사원의 시트 탭을 클릭")
    print("    2. Ctrl+P → 프린터: 'Microsoft Print to PDF'")
    print("    3. 파일명에 사원명 입력 후 저장")
    print("\n" + "="*60)

# ══════════════════════════════════════════════════════════
# 메인 실행
# ══════════════════════════════════════════════════════════
def main():
    print("\n" + "="*60)
    print("  판매수수료 자동계산 시스템")
    print("="*60)

    # 정산 월 입력
    while True:
        ym = input("\n  정산 월을 입력하세요 (예: 2026-05 또는 202605) : ").strip()
        # YYYYMM → YYYY-MM 자동 변환
        if len(ym) == 6 and ym.isdigit():
            ym = ym[:4] + "-" + ym[4:]
        try:
            datetime.strptime(ym, "%Y-%m")
            break
        except ValueError:
            print("  ❌ 형식이 맞지 않습니다. 2026-05 또는 202605 형식으로 입력해주세요.")

    year, month = ym.split("-")
    print(f"\n  ▶ {year}년 {int(month)}월 처리를 시작합니다...")

    # 파일 경로 확인
    master_path  = INPUT / "사원마스터.xlsx"

    ym_compact = ym.replace("-","")

    # 매장공제건집계 파일 탐지
    deduct_candidates = (
        list(INPUT.glob(f"*매장공제건*{ym[2:4]}년*.xlsx")) +
        list(INPUT.glob(f"*매장공제건*{ym[5:]}월*.xlsx")) +
        list(INPUT.glob(f"*매장공제건*.xlsx")) +
        list(INPUT.glob(f"◈*.xlsx"))
    )
    seen_d=set()
    deduct_candidates=[p for p in deduct_candidates
                       if not (str(p) in seen_d or seen_d.add(str(p)))]
    deduct_path = deduct_candidates[0] if deduct_candidates else None

    errors = []
    if not master_path.exists():
        errors.append(f"  ❌ 파일 없음: {master_path.name}")
    if errors:
        print("\n".join(errors))
        print("\n  input 폴더에 파일을 넣고 다시 실행하세요.")
        sys.exit(1)

    print("\n  📂 파일 읽는 중...")
    employees = load_master(master_path)
    deduction = load_deduction(deduct_path)
    print(f"     사원마스터: {len(employees)}명")
    if deduct_path:
        print(f"     공제건집계: {len(deduction)}개 매장  ({deduct_path.name})")
    else:
        print(f"     공제건집계: 없음 (공제 0 처리)")

    # 매출 파일(영*판매*.xlsx) 탐지 — calc·개별통지문·매출집계 모두에서 사용
    sales_candidates = (
        list(INPUT.glob(f"영*판매*{ym_compact}*.xlsx")) +
        list(INPUT.glob(f"영*{ym_compact}*판매*.xlsx")) +
        list(INPUT.glob(f"영*판매*{ym[2:4]}*.xlsx")) +
        list(INPUT.glob(f"영*{ym[5:]}*.xlsx"))
    )
    seen_paths = set()
    sales_candidates = [p for p in sales_candidates
                        if not (str(p) in seen_paths or seen_paths.add(str(p)))]

    # 영판매 파일에서 매출 데이터 로드 → calc()에 사용
    if sales_candidates:
        print(f"     영판매파일:  {sales_candidates[0].name}")
        try:
            from sales_report import load_sales_file as _lsf
            _sd_for_calc = _lsf(sales_candidates[0])
            sales = {shop: {"off": d["off_total"], "on": d["on_total"]}
                     for shop, d in _sd_for_calc.items()}
            print(f"     매출데이터:  {len(sales)}개 매장 로드")
        except Exception as _e:
            print(f"     ⚠️  매출 파일 읽기 오류: {_e}")
            sales = {}
    else:
        print(f"     영판매파일:  없음 (매출 0 처리)")
        sales = {}

    # 경비내역서 파일 탐지 (개별통지문·경비집계 모두에서 사용)
    exp_path = next(iter(INPUT.glob("★*경비내역서*.xlsx")), None) or \
               next(iter(INPUT.glob("*경비내역서*.xlsx")), None)
    from expense_report import load_expense

    # 계산
    print("\n  ⚙️  수수료 계산 중...")
    results = []
    missing_sales = []

    for emp in employees:
        if not emp["name"]:
            continue
        if emp["shop"] not in sales:
            missing_sales.append(emp["shop"])
        res = calc(emp, sales, deduction)
        results.append({"emp": emp, "res": res})

    missing_sales = list(set(missing_sales))
    print(f"     계산 완료: {len(results)}명")

    # 출력 폴더 생성
    out_dir = OUTPUT / ym
    out_dir.mkdir(parents=True, exist_ok=True)

    # 파일 생성
    print("\n  📝 파일 생성 중...")
    # 은행이체목록.xlsx 생성 제거 (불필요)
    # 매출 상세 + 경비 데이터 → 개별통지문에 전달
    sales_detail_for_notice = {}
    expense_rows_for_notice = {}
    if sales_candidates:
        try:
            from sales_report import load_sales_file
            _sd = load_sales_file(sales_candidates[0])
            sales_detail_for_notice = _sd
        except Exception:
            pass
    # 경비 행별 데이터 구성
    try:
        from expense_report import (load_expense as _le, load_deduction_detail as _ldd,
                                    load_inventory as _linv, STORE_CODE_MAP as _scm)
        _exp_d = _le(exp_path)
        _ded_d = _ldd(deduct_path)
        _inv_path2 = next(iter(INPUT.glob("*재고실사*공제건*합계*.xlsx")),None) or                      next(iter(INPUT.glob("*재고실사*.xlsx")),None)
        _inv_d = _linv(_inv_path2) if _inv_path2 else {}
        _event2 = {e["shop"] for e in employees if (e.get("note","") or "").strip()=="행사매장"}
        _pos_t2 = sum(d.get("pos",0) for d in _ded_d.values())
        _pos_z2 = sum(1 for s,d in _ded_d.items() if d.get("pos",0)==0 and s not in _event2)
        _pu2 = int(_pos_t2/_pos_z2/10)*10 if _pos_z2 else 0
        for _shop in _scm.keys():
            _d2=_ded_d.get(_shop,{})
            expense_rows_for_notice[_shop] = {
                "j_direct":  _exp_d.get(_shop,0),
                "shortfall": _d2.get("shortfall",0),
                "loss":      _inv_d.get(_shop,0),
                "gift":      _d2.get("gift",0),
                "pos":       _d2.get("pos",0),
                "t_refund":  _pu2 if _d2.get("pos",0)==0 and _shop not in _event2 else 0,
            }
    except Exception as _e2:
        print(f"     ⚠️  경비 데이터 로드 오류: {_e2}")
    # ── 아르바이트 정산서 처리 ──────────────────────────────
    print("\n  🧾 아르바이트 정산서 처리 중...")
    ym_compact = ym.replace("-", "")

    # ── 아르바이트 파일: 파일명에 '일용직' 포함 여부 무관하게 glob 검색
    def _find(pattern_list):
        for pat in pattern_list:
            found = list(INPUT.glob(pat))
            if found: return found[0]
        return None

    arba_staff  = _find([f"사원현황_{ym_compact}.xlsx",  f"사원현황*{ym_compact}*.xlsx",  f"사원현황*일용직*{ym_compact}*.xlsx"])
    arba_salary = _find([f"급여상여명세서_{ym_compact}.xlsx", f"급여상여명세서*{ym_compact}*.xlsx", f"급여상여명세서*일용직*{ym_compact}*.xlsx"])
    arba_attend = _find([f"월별근태_{ym_compact}.xlsx",   f"월별근태*{ym_compact}*.xlsx",   f"월별근태*일용직*{ym_compact}*.xlsx"])

    arba_files_exist = all([arba_staff, arba_salary, arba_attend])
    if arba_files_exist:
        try:
            from arba_process import run as arba_run
            arba_run(ym, BASE)
            # (아르바이트 print는 arba_process.py 내부에서 출력)
        except Exception as e:
            print(f"     ⚠️  아르바이트 처리 중 오류: {e}")
    else:
        missing = []
        if not arba_staff:  missing.append(f"사원현황_{ym_compact}.xlsx  (또는 사원현황_일용직_{ym_compact}.xlsx)")
        if not arba_salary: missing.append(f"급여상여명세서_{ym_compact}.xlsx  (또는 급여상여명세서_일용직_{ym_compact}.xlsx)")
        if not arba_attend: missing.append(f"월별근태_{ym_compact}.xlsx  (또는 월별근태_일용직_{ym_compact}.xlsx)")
        print(f"     ℹ️  아르바이트 파일 없음 → 건너뜀")
        for m in missing: print(f"        - {m}")

    # ── 매출집계 보고서 처리 ──────────────────────────────────
    print("\n  📊 매출집계 보고서 처리 중...")
    if not sales_candidates:
        print(f"     ℹ️  매출 파일 없음 → 건너뜀")
        print(f"        (필요 시 input 폴더에 추가: 영YY-MM_판매_{ym_compact}.xlsx 또는 영5-9_판매_{ym_compact}.xlsx)")
    else:
        try:
            from sales_report import load_sales_file, make_sales_report, STORE_CODE_MAP
            # 사원마스터에서 매니저명 + 소득구분 추출
            mgr_map    = {}
            income_map = {}
            for emp in employees:
                shop = emp.get("shop", "")
                if not shop or not emp.get("name"):
                    continue
                pay_type = emp.get("pay_type", "")
                income   = emp.get("income", "")
                # 매장 대표자(매니저 or 본사-M) 우선
                if pay_type == "중간관리":
                    mgr_map[shop]    = emp["name"]
                    income_map[shop] = "중간관리"
                elif emp.get("grade", "") == "본사-M" and shop not in mgr_map:
                    mgr_map[shop]    = emp["name"]
                    income_map[shop] = income  # 근로소득 or 사업소득
            sales_data = load_sales_file(sales_candidates[0])
            make_sales_report(sales_data, ym, out_dir, mgr_map, income_map)
        except Exception as e:
            print(f"     ⚠️  매출집계 처리 중 오류: {e}")

    # ── 경비지원및공제 집계 처리 ─────────────────────────────
    try:
        from expense_report import run as expense_run
        expense_run(ym, BASE, employees)
    except Exception as e:
        print(f"     ⚠️  경비지원및공제 처리 오류: {e}")

    # ── 매출집계(분석) 보고서 처리 ───────────────────────────
    # (매출집계분석 print는 sales_analysis.py 내부에서 출력)
    _pu_val = 0   # POS 단위환급액 기본값 (try 실패 시에도 안전하게 사용)

    if sales_candidates:
        try:
            from sales_analysis import run as analysis_run, load_sales_vp_cg
            from expense_report import (load_deduction_detail, load_inventory,
                                        STORE_CODE_MAP as _SCM)
            # 경비합계 재계산 (expense_report와 동일 로직)
            _ded = load_deduction_detail(deduct_path)
            _inv_path = next(iter(INPUT.glob("*재고실사*공제건*합계*.xlsx")), None) or \
                        next(iter(INPUT.glob("*재고실사*.xlsx")), None)
            from expense_report import load_inventory as _li
            _inv = _li(_inv_path) if _inv_path else {}
            _event = {e["shop"] for e in employees
                      if (e.get("note","") or "").strip()=="행사매장"}
            _pos_t = sum(d.get("pos",0) for d in _ded.values())
            _pos_z = sum(1 for s,d in _ded.items()
                         if d.get("pos",0)==0 and s not in _event)
            _pu = int(_pos_t/_pos_z/10)*10 if _pos_z else 0
            _pu_val = _pu   # try 성공 시에만 갱신
            _exp_map = {}
            for _shop in _SCM.keys():
                _d=_ded.get(_shop,{}); _j=load_expense(exp_path).get(_shop,0) if exp_path else 0
                _l=_d.get("shortfall",0); _n=_inv.get(_shop,0)
                _o=_d.get("gift",0); _s=_d.get("pos",0)
                _t=_pu if _s==0 and _shop not in _event else 0
                _exp_map[_shop]=_j-_l-_n-_o-_s+_t
            # 아르바이트_정산서 경로 (이미 생성된 output 파일)
            _arba_path = out_dir / "아르바이트_정산서.xlsx"
            analysis_run(ym, BASE, results, _exp_map,
                         sales_candidates[0], employees,
                         _arba_path if _arba_path.exists() else None)
        except Exception as e:
            print(f"     ⚠️  매출집계(분석) 처리 오류: {e}")
    else:
        print(f"     ℹ️  영판매 파일 없음 → 건너뜀")

    # ── 판매수수료작업시트(AMD) 처리 ────────────────────────────
    # (AMD print는 amd_report.py 내부에서 출력)
    try:
        from amd_report import run as amd_run
        amd_run(ym, BASE, results, sales_detail_for_notice,
                expense_rows_for_notice, employees)
    except Exception as e:
        print(f"     ⚠️  AMD 처리 오류: {e}")
        import traceback; traceback.print_exc()

    # ── 개별통지문 처리 (AMD 이후) ──────────────────────────────
    biz_cnt = sum(1 for r in results if r["emp"]["income"]=="사업소득")
    p3 = make_notice(results, ym, out_dir, sales_detail_for_notice, expense_rows_for_notice)
    print(f"     ✅ 개별통지문.xlsx  (사업소득자 {biz_cnt}명)")

    # ── 판매수수료집계 처리 ───────────────────────────────────
    # (판매수수료집계 print는 sales_summary.py 내부에서 출력)
    try:
        from sales_summary import run as summary_run
        _ev = {e["shop"] for e in employees
               if (e.get("note","") or "").strip()=="행사매장"}
        # _pu_val은 위 try 블록에서 이미 설정됨
        summary_run(ym, BASE, results,
                    sales_detail_for_notice,
                    expense_rows_for_notice,
                    employees,
                    pos_unit=_pu_val,
                    event_shops=_ev)
    except Exception as e:
        print(f"     ⚠️  판매수수료집계 오류: {e}")
        import traceback; traceback.print_exc()

    # ── 중간관리판매수수료이체내역 처리 ───────────────────────
    try:
        from transfer_report import run as transfer_run
        transfer_run(ym, BASE, results,
                     expense_rows_for_notice, employees)
    except Exception as e:
        print(f"     ⚠️  이체내역 처리 오류: {e}")
        import traceback; traceback.print_exc()

    # ── 매입세금계산서 처리 ────────────────────────────────────
    try:
        from tax_invoice_report import run as tax_run
        tax_run(ym, BASE, results,
                expense_rows_for_notice, employees)
    except Exception as e:
        print(f"     ⚠️  세금계산서 처리 오류: {e}")
        import traceback; traceback.print_exc()

    # ── 처리 결과 요약 (아르바이트 포함 모든 처리 완료 후) ──
    p4 = make_summary(results, ym, out_dir, missing_sales,
                      expense_rows=expense_rows_for_notice, master=employees)


    # ── 월별손익DB 업데이트 및 연도/분기/시즌 보고서 생성 ─────────────
    try:
        from history_update import run as history_run
        history_run(ym, BASE)
    except Exception as e:
        print(f"     ⚠️  월별손익DB 업데이트 오류: {e}")
        import traceback; traceback.print_exc()

    # PDF 안내
    print_pdf_guide(p3)

    # 출력 폴더 열기 (Windows)
    try:
        os.startfile(str(out_dir))
    except Exception:
        pass

    print(f"\n  ✅ 모든 처리가 완료되었습니다!")
    print(f"     출력 폴더: {out_dir}")

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n  ❌ 예상치 못한 오류가 발생했습니다: {e}")
        import traceback
        traceback.print_exc()
    finally:
        pass
